"""Actual company-workbook calibration photos through the offline OCR engine.

Both PNG fixtures are extracted byte for byte from the supplied XLSX and their
real drawing anchors are checked. No printed-label fixture, OCR response,
recognized word, confidence, or identity is fabricated. image45 is a date-only
label, so this test does not claim it independently identifies an engineer.
image37 contains the actual serial barcode and can exercise automatic matching.
The original workbook already contains these photos' dates. A temporary copy
sets just the serial-bearing photo's two date cells one year earlier, preserving
all other ZIP entries. This ensures automatic writing uses the recognized dates
instead of copying the pre-existing values. Neither original input is changed.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import posixpath
from pathlib import Path
import re
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from PIL import Image, ImageChops, ImageStat
from playwright.sync_api import sync_playwright

from browser_test import REPO, serve_program


CASES = [
    {"media": "xl/media/image45.png", "sheet": "광주", "row": 6, "identity_visible": False},
    {"media": "xl/media/image37.png", "sheet": "광주", "row": 2, "identity_visible": True},
]


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--diagnose-only", action="store_true", help="Record a previous version's real OCR without requiring corrected dates or automatic writing.")
    parser.add_argument("--diagnostics", type=Path, default=Path("/tmp/tool-calibration-label-ocr.json"))
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    if args.workbook is None or not args.workbook.is_file():
        parser.error("Supply --workbook with the attached company XLSX.")
    if not args.program.is_file():
        parser.error("Supply --program with the standalone HTML program.")
    return args


def extract_original_cases(workbook, temporary):
    source = load_workbook(workbook)
    cases = []
    with ZipFile(workbook) as archive:
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        relationship_targets = {item.attrib["Id"]: item.attrib["Target"] for item in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        sheet_paths = {}
        for item in ET.fromstring(archive.read("xl/workbook.xml")).findall("{" + namespace + "}sheets/{" + namespace + "}sheet"):
            target = relationship_targets[item.attrib["{" + relationship_namespace + "}id"]]
            sheet_paths[item.attrib["name"]] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        for expected in CASES:
            case = dict(expected)
            data = archive.read(case["media"])
            sheet = source[case["sheet"]]
            headers = {str(cell.value): cell.column for cell in sheet[1]}
            # Expected values come from the associated row, not company data
            # copied into this test source. These two original labels were also
            # visually checked against their already-recorded worksheet dates.
            case["serial"] = str(sheet.cell(case["row"], headers["Serial"]).value)
            case["name"] = str(sheet.cell(case["row"], headers["Asset #"]).value).rsplit("_", 1)[-1].strip()
            case["last"] = sheet.cell(case["row"], headers["Cal Last"]).value.strftime("%Y-%m-%d")
            case["due"] = sheet.cell(case["row"], headers["Cal Next"]).value.strftime("%Y-%m-%d")
            case["sheet_path"] = sheet_paths[case["sheet"]]
            case["date_cells"] = {"last": get_column_letter(headers["Cal Last"]) + str(case["row"]), "due": get_column_letter(headers["Cal Next"]) + str(case["row"])}
            assert case["serial"] and case["name"]
            anchored = [image for image in sheet._images if image.anchor._from.row == case["row"] - 1 and image.anchor._from.col == headers["Label Photo"] - 1]
            assert len(anchored) == 1 and sha256(anchored[0]._data()).digest() == sha256(data).digest(), case
            with Image.open(BytesIO(data)) as image:
                case["dimensions"] = list(image.size)
            case["sha256"] = sha256(data).hexdigest()
            case["path"] = temporary / Path(case["media"]).name
            case["path"].write_bytes(data)
            assert case["path"].read_bytes() == archive.read(case["media"])
            cases.append(case)
    source.close()
    return cases


def past_date_fixture(original, temporary, cases):
    """Patch only two numeric cell values in a disposable copy of the XLSX."""
    book = load_workbook(original, read_only=True)
    epoch = book.epoch
    book.close()
    patches = {}
    for case in cases:
        if not case["identity_visible"]:
            continue
        case["prior_dates"] = {}
        for role, cell in case["date_cells"].items():
            current = datetime.fromisoformat(case[role])
            try:
                prior = current.replace(year=current.year - 1)
            except ValueError:
                prior = current.replace(year=current.year - 1, day=28)
            case["prior_dates"][role] = prior.strftime("%Y-%m-%d")
            number = to_excel(prior, epoch)
            patches.setdefault(case["sheet_path"], {})[cell] = str(int(number)) if number.is_integer() else str(number)
    output = temporary / "past-date-workbook.xlsx"
    with ZipFile(original) as before, ZipFile(output, "w") as after:
        for info in before.infolist():
            data = before.read(info.filename)
            if info.filename in patches:
                text = data.decode("utf-8")
                for cell, value in patches[info.filename].items():
                    pattern = r'(<c\b[^>]*\br="' + re.escape(cell) + r'"[^>]*>.*?<v>)([^<]*)(</v>)'
                    text, count = re.subn(pattern, lambda match: match[1] + value + match[3], text, flags=re.S)
                    assert count == 1, (info.filename, cell, count)
                data = text.encode("utf-8")
            after.writestr(info, data)
    with ZipFile(original) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
        assert changed == set(patches), changed
    return output


def verify_actual_photo_workbook(original, updated, case):
    before, after = load_workbook(original), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    checked = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for row in old_sheet:
            for old_cell in row:
                new_cell = new_sheet[old_cell.coordinate]
                is_target_date = old_sheet.title == case["sheet"] and old_cell.coordinate in case["date_cells"].values()
                if not is_target_date:
                    assert old_cell.value == new_cell.value, (old_sheet.title, old_cell.coordinate)
                assert old_cell._style == new_cell._style, (old_sheet.title, old_cell.coordinate)
                checked += 1
    assert sum(len(sheet._images) for sheet in before) == sum(len(sheet._images) for sheet in after) == 59
    sheet = after[case["sheet"]]
    headers = {str(cell.value): cell.column for cell in sheet[1]}
    assert sheet.cell(case["row"], headers["Cal Last"]).value.strftime("%Y-%m-%d") == case["last"]
    assert sheet.cell(case["row"], headers["Cal Next"]).value.strftime("%Y-%m-%d") == case["due"]
    allowed = {"[Content_Types].xml", "xl/worksheets/sheet2.xml", "xl/drawings/drawing2.xml", "xl/drawings/_rels/drawing2.xml.rels"}
    with ZipFile(original) as old_zip, ZipFile(updated) as new_zip:
        old_names, new_names = set(old_zip.namelist()), set(new_zip.namelist())
        assert old_names <= new_names
        changed = [name for name in old_names if old_zip.read(name) != new_zip.read(name)]
        assert not set(changed) - allowed, changed
        added = [name for name in new_names - old_names if not name.endswith("/")]
        assert len(added) == 1 and added[0].startswith("xl/media/"), added
        data = new_zip.read(added[0])
        anchored = [image for image in sheet._images if image.anchor._from.row == case["row"] - 1 and image.anchor._from.col == headers["Label Photo"] - 1]
        assert len(anchored) == 1 and sha256(anchored[0]._data()).digest() == sha256(data).digest()
        with Image.open(BytesIO(data)) as inserted, Image.open(case["path"]) as original_photo:
            assert inserted.size == original_photo.size
            difference = ImageStat.Stat(ImageChops.difference(inserted.convert("RGB"), original_photo.convert("RGB")))
            assert max(difference.mean) < 5, difference.mean
    assert before[case["sheet"]][case["date_cells"]["last"]].value.strftime("%Y-%m-%d") == case["prior_dates"]["last"]
    assert before[case["sheet"]][case["date_cells"]["due"]].value.strftime("%Y-%m-%d") == case["prior_dates"]["due"]
    before.close()
    after.close()
    return {"cells_checked": checked, "pictures": 59, "correct_photo_anchor": True, "actual_photo_content_matches": True, "prior_dates": case["prior_dates"], "written_dates": {"last": case["last"], "due": case["due"]}}


def main():
    args = arguments()
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    diagnostics = {"program": str(args.program), "diagnose_only": args.diagnose_only, "actual_workbook_photos": True, "synthetic_photo": False, "cases": []}
    failures = []
    with tempfile.TemporaryDirectory(prefix="tool-real-label-ocr-") as directory:
        temporary = Path(directory)
        cases = extract_original_cases(args.workbook, temporary)
        input_workbook = past_date_fixture(args.workbook, temporary, cases)
        with serve_program(args.program) as url, sync_playwright() as p:
            browser = p.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
            context = browser.new_context(accept_downloads=True, locale="ko-KR")
            page = context.new_page()
            errors, requests = [], []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(url)
            page.wait_for_function("window.ToolApp&&window.ToolOCR&&window.ToolRuntime")
            # Observe the real worker's unmodified result before the parser
            # discards details. This wrapper never replaces a response.
            page.evaluate("""() => {
              window.calibrationRawReadings=[];
              const create=Tesseract.createWorker.bind(Tesseract);
              Tesseract.createWorker=async(...args)=>{
                const worker=await create(...args),recognize=worker.recognize.bind(worker);
                worker.recognize=async(...input)=>{
                  const result=await recognize(...input),data=result.data||{};
                  const words=(data.words||[]).map(word=>({text:word.text,
                    confidence:word.confidence,bbox:word.bbox,
                    symbols:(word.symbols||[]).map(symbol=>({text:symbol.text,
                      confidence:symbol.confidence,bbox:symbol.bbox}))}));
                  calibrationRawReadings.push({text:data.text,confidence:data.confidence,words});
                  return result;
                };return worker;
              };
            }""")
            context.route("**/*", lambda route: (requests.append(route.request.url), route.abort()) if route.request.url.startswith(("http:", "https:")) else route.continue_())
            context.set_offline(True)
            page.locator("#workbook").set_input_files(input_workbook)
            page.wait_for_function("ToolApp.state.model&&!ToolApp.state.busy")
            for case in cases:
                page.evaluate("calibrationRawReadings=[]")
                page.locator("#photos").set_input_files(case["path"])
                page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing", timeout=180000)
                result = page.evaluate("""() => {
                  const photo=ToolApp.state.photos[0];return {
                    result:photo.result,approval:photo.approval,assessment:photo.assessment,
                    prepared:!!ToolApp.state.preparedResult,readings:calibrationRawReadings,
                    fields:{last:document.querySelector('#last').value,due:document.querySelector('#due').value}
                  };
                }""")
                report = {key: value for key, value in case.items() if key != "path"}
                report.update(result)
                diagnostics["cases"].append(report)
                args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
                print(json.dumps({key: value for key, value in report.items() if key not in {"readings", "result"}}, ensure_ascii=False, indent=2), flush=True)
                if not args.diagnose_only:
                    recognized = result["result"] or {}
                    if recognized.get("failed") or (recognized.get("last"), recognized.get("due")) != (case["last"], case["due"]):
                        failures.append(case["media"] + ": actual Korean calibration dates were not read correctly")
                    if case["identity_visible"]:
                        if not result["prepared"] or result["approval"] != "automatic" or not (result["assessment"] or {}).get("eligible"):
                            failures.append(case["media"] + ": actual serial-bearing photo did not prepare an automatic workbook")
                        else:
                            assert result["assessment"]["serial"] == case["serial"]
                            assert result["assessment"]["name"] == case["name"]
                            assert result["assessment"]["toolId"] == case["sheet_path"] + ":" + str(case["row"])
                            with page.expect_download() as event:
                                page.locator("#automatic-result-download").click()
                            updated = temporary / "actual-label-automatic-result.xlsx"
                            event.value.save_as(updated)
                            report["workbook"] = verify_actual_photo_workbook(input_workbook, updated, case)
                    else:
                        assert not result["prepared"] and not (result["assessment"] or {}).get("eligible"), "A date-only label cannot independently identify an engineer"
                page.locator("#remove-photo").click()
                page.wait_for_function("ToolApp.state.photos.length===0&&!ToolApp.state.preparing")
            assert not requests, requests
            assert not errors, errors
            browser.close()
    assert original_hash == sha256(args.workbook.read_bytes()).hexdigest(), "Original workbook changed"
    diagnostics.update({"failures": failures, "http_requests": requests, "page_errors": errors, "original_unchanged": True})
    args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"actual_photo_cases": len(cases), "diagnose_only": args.diagnose_only, "diagnostics": str(args.diagnostics), "failures": failures, "http_requests": requests, "original_unchanged": True}, ensure_ascii=False, indent=2))
    assert not failures, failures


if __name__ == "__main__":
    main()
