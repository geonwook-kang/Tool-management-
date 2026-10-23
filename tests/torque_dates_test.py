"""Torque-label format and XLSX mapping regressions, using controlled OCR text.

The reported screenshot names serial 0716805829, calibration 2026.06.22 and
next calibration 2026.12.22. Its original JPEG is unavailable to this test.
These checks cover parser decisions, split OCR words, exact leading-zero
identity, and the generated Excel. They do not measure screenshot OCR accuracy.
The uploaded company workbook is read only and generated files stay temporary.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import tempfile
from urllib.parse import urlsplit

from openpyxl import load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright

from browser_test import REPO, arguments, serve_program


SERIAL = "0716805829"
LAST, DUE = "2026-06-22", "2026-12-22"
CERTIFICATE = "DA26-0554-002"
LABELS = "교정일자 2026.06.22\n차기교정예정일자 2026.12.22"
TEXT = SERIAL + "\n성적서번호 " + CERTIFICATE + "\n" + LABELS


def positioned_words(separator=".", date_confidence=95):
    words = [
        {"text": SERIAL, "confidence": 95, "bbox": {"x0": 20, "y0": 10, "x1": 290, "y1": 40}},
        {"text": "성적서번호", "confidence": 95, "bbox": {"x0": 20, "y0": 55, "x1": 150, "y1": 80}},
        {"text": CERTIFICATE, "confidence": 95, "bbox": {"x0": 310, "y0": 55, "x1": 600, "y1": 80}},
        {"text": "교정일자", "confidence": 95, "bbox": {"x0": 20, "y0": 110, "x1": 230, "y1": 140}},
        {"text": "차기교정예정일자", "confidence": 95, "bbox": {"x0": 20, "y0": 190, "x1": 230, "y1": 220}},
    ]
    for y, month in [(110, "06"), (190, "12")]:
        for token, x0, x1, y0, y1 in [
            ("2026", 310, 400, y, y + 30),
            (separator, 408, 415, y + 23, y + 29),
            (month, 430, 470, y, y + 30),
            (separator, 478, 485, y + 23, y + 29),
            ("22", 500, 540, y, y + 30),
        ]:
            words.append({"text": token, "confidence": date_confidence,
                          "bbox": {"x0": x0, "y0": y0, "x1": x1, "y1": y1}})
    return words


def captured_refinement_passes():
    """Minimal date-only capture from the actual candidate-3 label diagnosis.

    Coordinates and word confidence come from its two refinement responses.
    The low-confidence 29 belongs to a different OCR line, while its own box
    overlaps the true June 24 row. It must not create an unobserved June 29.
    """
    last_line = {"x0": 26, "y0": 12, "x1": 217, "y1": 44}
    due_line = {"x0": 27, "y0": 59, "x1": 219, "y1": 94}
    junk_line = {"x0": 50, "y0": 94, "x1": 221, "y1": 107}
    def word(text, confidence, coordinates, line):
        return {"text": text, "confidence": confidence,
                "bbox": dict(zip(["x0", "y0", "x1", "y1"], coordinates)), "lineBox": line}
    last_words = [
        word("2025.", 95.15010833740234, (26, 17, 121, 44), last_line),
        word("12.24", 74.71382141113281, (138, 12, 217, 40), last_line),
    ]
    return [
        {"source": "date-crop-refine-6", "confidence": 61,
         "text": "2025. 12.24\n2026. 06. 24\n4240. 29ᆞ77\n", "words": last_words + [
             word("2026.", 82.14422607421875, (27, 66, 123, 94), due_line),
             word("06.", 82.14422607421875, (138, 59, 219, 89), due_line),
             word("24", 92.2001724243164, (195, 55, 219, 103), due_line),
             word("4240.", 69.75782775878906, (50, 94, 221, 107), junk_line),
             word("29", 6.4109954833984375, (138, 71, 177, 107), junk_line),
             word("ᆞ", 52.086673736572266, (176, 71, 188, 107), junk_line),
             word("77", 0, (187, 71, 220, 107), junk_line),
         ]},
        {"source": "date-crop-refine-11", "confidence": 80,
         "text": "2025. 12.24\n\n2026. 06. 24\n", "words": last_words + [
             word("2026.", 72.91075134277344, (27, 66, 123, 94), due_line),
             word("06.", 72.91075134277344, (138, 59, 219, 89), due_line),
             word("24", 90.5966796875, (194, 55, 219, 104), due_line),
         ]},
    ]


def formats():
    split = SERIAL + "\n성적서번호 " + CERTIFICATE + "\n교정일자\n차기교정예정일자\n2026 . 06 . 22\n2026 . 12 . 22"
    comma = SERIAL + "\n교정일자 2026,06,22\n차기교정예정일자 2026,12,22"
    owner_text = "__OWNER__\n" + LABELS
    probe = {"text": "Serial: 999999999\n2026.08.24", "confidence": 80, "source": "date-probe"}
    clean_refine = "2026.06.22\n2026.12.22"
    return {
        "reported_korean_labels_and_dates": {"text": TEXT, "eligible": True},
        "spaces_between_date_components": {"text": SERIAL + "\n교정일자 2026 . 06 . 22\n차기교정예정일자 2026 . 12 . 22", "eligible": True},
        "separate_date_words_and_low_position_dots": {"text": split, "words": positioned_words(), "eligible": True},
        "reversed_ocr_word_order_keeps_physical_rows": {"text": split, "words": list(reversed(positioned_words())), "eligible": True},
        "unsupported_comma_dates_are_not_guessed": {"text": comma, "eligible": False},
        "low_confidence_split_date_words_need_review": {"text": split, "words": positioned_words(date_confidence=25), "eligible": False},
        "certificate_is_not_tool_identity": {"text": "성적서번호 " + CERTIFICATE + "\n" + LABELS, "eligible": False},
        "missing_leading_zero_is_not_exact_serial": {"text": "716805829\n" + LABELS, "eligible": False},
        "letter_o_is_not_leading_zero": {"text": "O716805829\n" + LABELS, "eligible": False},
        "impossible_last_day_is_not_repaired": {"text": SERIAL + "\n교정일자 2026.06.32\n차기교정예정일자 2026.12.22", "eligible": False},
        "impossible_due_day_is_not_repaired": {"text": SERIAL + "\n교정일자 2026.06.22\n차기교정예정일자 2026.12.32", "eligible": False},
        "manufacture_date_does_not_replace_calibration": {"text": SERIAL + "\n제조일자 2026.01.01\n성적서번호 " + CERTIFICATE + "\n" + LABELS, "eligible": True},
        "date_discovery_probe_cannot_supply_approval_dates": {"text": SERIAL, "passes": [
            {"text": SERIAL, "confidence": 95, "source": "original"}, probe], "eligible": False},
        "probe_and_one_clean_unlabelled_read_are_not_agreement": {"text": SERIAL, "passes": [
            {"text": SERIAL, "confidence": 95, "source": "original"}, probe,
            {"text": clean_refine, "confidence": 95, "source": "date-crop-refine-6"}], "eligible": False},
        "two_clean_refinements_supply_dates_independently_of_probe": {"text": SERIAL, "passes": [
            {"text": SERIAL, "confidence": 95, "source": "original"}, probe,
            {"text": clean_refine, "confidence": 95, "source": "date-crop-refine-6"},
            {"text": clean_refine, "confidence": 95, "source": "date-crop-refine-11"}],
            "eligible": True, "dateMethod": "repeated-dates"},
        "actual_capture_separate_line_junk_cannot_manufacture_june_29": {"text": SERIAL, "passes": [
            {"text": SERIAL, "confidence": 95, "source": "original"}] + captured_refinement_passes(),
            "eligible": True, "dateMethod": "repeated-dates", "last": "2025-12-24", "due": "2026-06-24",
            "expectedDates": ["2025-12-24", "2026-06-24"]},
        "genuine_high_confidence_june_29_conflict_remains_pending": {"text": SERIAL, "passes": [
            {"text": SERIAL + "\n교정일자 2025.12.24\n차기교정예정일자 2026.06.24", "confidence": 95, "source": "original"},
            {"text": "교정일자 2025.12.24\n차기교정예정일자 2026.06.29", "confidence": 95, "source": "date-crop"}],
            "eligible": False, "expectedDates": ["2025-12-24", "2026-06-24", "2026-06-29"]},
        "name_cycle_selects_unique_six_month_tool": {"text": owner_text, "eligible": True, "matchMethod": "name-cycle"},
        "name_cycle_selects_unique_twelve_month_tool": {"text": "__OWNER__\n교정일자 2026.06.22\n차기교정예정일자 2027.06.22", "eligible": True,
                                                        "matchMethod": "name-cycle", "expectedMonths": 12, "due": "2027-06-22"},
        "name_cycle_preserves_current_excel_eighteen_month_interval": {"text": "__OWNER__\n교정일자 2026.06.22\n차기교정예정일자 2027.12.22", "eligible": True,
                                                                      "matchMethod": "name-cycle", "expectedMonths": 18,
                                                                      "replaceTwelveMonthInterval": 18, "due": "2027-12-22"},
        "name_cycle_same_owner_same_interval_stays_ambiguous": {"text": owner_text, "eligible": False, "sameOwnerSixMonthTools": True},
        "name_cycle_unknown_explicit_serial_blocks_fallback": {"text": "__OWNER__\nSerial: 999999999\n" + LABELS, "eligible": False},
        "name_cycle_low_confidence_dates_do_not_resolve_identity": {"text": owner_text, "words": [
            {"text": "__OWNER__", "confidence": 95}, {"text": "교정일자", "confidence": 95},
            {"text": "2026.06.22", "confidence": 25}, {"text": "차기교정예정일자", "confidence": 95},
            {"text": "2026.12.22", "confidence": 25}], "eligible": False},
        "name_cycle_conflicting_dates_do_not_resolve_identity": {"text": owner_text, "passes": [
            {"text": owner_text, "confidence": 95, "source": "original"},
            {"text": "교정일자 2026.06.22\n차기교정예정일자 2026.12.21", "confidence": 95, "source": "date-crop"}], "eligible": False},
    }


def policy_checks(page, url, workbook):
    address = urlsplit(url)
    page.route("**/torque-label-policy.html", lambda route: route.fulfill(
        status=200, content_type="text/html", body='<!doctype html><input id="fixture" type="file">'))
    page.goto(f"{address.scheme}://{address.netloc}/torque-label-policy.html")
    for source in ["vendor/jszip.min.js", "workbook.js", "ocr.js", "automation.js"]:
        page.add_script_tag(path=REPO / "web" / source)
    page.locator("#fixture").set_input_files(workbook)
    outcomes = page.evaluate("""async ({cases,serial}) => {
        const model=await ToolWorkbook.load(await document.querySelector('#fixture').files[0].arrayBuffer());
        const target=model.tools.find(tool=>tool.serial===serial);
        if(!target)throw new Error('Exact leading-zero torque serial missing in company fixture');
        const result={};
        for(const [name,spec] of Object.entries(cases)){
            const confidence=spec.confidence??95;
            const fixtureText=text=>text.replaceAll('__OWNER__',target.name);
            const caseModel={...model,tools:model.tools.map(tool=>{
                if(tool.name!==target.name)return tool;
                const months=spec.sameOwnerSixMonthTools?6:
                    spec.replaceTwelveMonthInterval&&tool.months===12?spec.replaceTwelveMonthInterval:null;
                return months?{...tool,months,interval:String(months),cycle:'MONTHS'}:tool;
            })};
            const text=fixtureText(spec.text);
            const passes=(spec.passes||[{text,confidence,source:'original',words:spec.words||[]}]).map(pass=>({
                ...pass,text:fixtureText(pass.text),words:(pass.words||[]).map(word=>({...word,text:fixtureText(word.text)}))}));
            const parsed=ToolOCR.parse(text,caseModel.tools,{confidence,passes});
            const expected=spec.expectedMonths?caseModel.tools.find(tool=>tool.name===target.name&&tool.months===spec.expectedMonths):target;
            result[name]={assessment:ToolAutomation.assess(caseModel,parsed),
                candidates:parsed.candidateIds,observedSerials:parsed.identityEvidence.observedSerials,
                serialEvidence:parsed.identityEvidence.observedSerialEvidence,
                dates:parsed.dateEvidence,target:{id:target.id,row:target.row,serial:target.serial,months:target.months},
                expected:expected&&{id:expected.id,serial:expected.serial}};
        }
        return result;
    }""", {"cases": formats(), "serial": SERIAL})
    for name, spec in formats().items():
        outcome = outcomes[name]
        assessment = outcome["assessment"]
        assert assessment["eligible"] is spec["eligible"], (name, outcome)
        assert outcome["target"]["serial"] == SERIAL
        assert outcome["target"]["months"] == 6
        assert CERTIFICATE.replace("-", "") not in outcome["observedSerials"], (name, outcome)
        if spec["eligible"]:
            assert assessment["toolId"] == outcome["expected"]["id"], (name, outcome)
            assert assessment["serial"] == outcome["expected"]["serial"], (name, outcome)
            assert assessment["last"] == spec.get("last", LAST), (name, outcome)
            assert assessment["due"] == spec.get("due", DUE), (name, outcome)
            assert assessment["dateMethod"] == spec.get("dateMethod", "labelled"), (name, outcome)
            if spec.get("matchMethod"):
                assert assessment["matchMethod"] == spec["matchMethod"], (name, outcome)
        else:
            assert assessment["reasons"], (name, outcome)
        if spec.get("expectedDates"):
            assert outcome["dates"]["dates"] == spec["expectedDates"], (name, outcome)
    for case in ["certificate_is_not_tool_identity", "missing_leading_zero_is_not_exact_serial", "letter_o_is_not_leading_zero"]:
        assert outcomes[case]["candidates"] == [], (case, outcomes[case])
    return {"accepted_formats": [name for name, spec in formats().items() if spec["eligible"]],
            "pending_formats": [name for name, spec in formats().items() if not spec["eligible"]],
            "serial_remains_exact_string": SERIAL, "certificate_excluded_from_serial_evidence": True}


def verify_torque_workbook(original, updated):
    before, after = load_workbook(original), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    allowed = set()
    target = None
    for sheet in after:
        headings = {cell.value: cell.column for cell in sheet[1]}
        if "Serial" not in headings:
            continue
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headings["Serial"]).value == SERIAL:
                assert target is None, "Duplicate leading-zero fixture serial"
                target = {"sheet": sheet.title, "row": row}
                assert sheet.cell(row, headings["Serial"]).data_type == "s"
                for heading, date in [("Cal Last", LAST), ("Cal Next", DUE)]:
                    cell = sheet.cell(row, headings[heading])
                    assert cell.value == datetime.fromisoformat(date)
                    allowed.add((sheet.title, cell.coordinate))
    assert target is not None
    cells = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for row in old_sheet:
            for old in row:
                new = new_sheet[old.coordinate]
                if (old_sheet.title, old.coordinate) not in allowed:
                    assert old.value == new.value, (old_sheet.title, old.coordinate)
                assert old._style == new._style, (old_sheet.title, old.coordinate)
                cells += 1
    pictures = sum(len(sheet._images) for sheet in after)
    assert pictures == sum(len(sheet._images) for sheet in before) == 59
    before.close()
    after.close()
    return {"target": target, "Cal Last": LAST, "Cal Next": DUE,
            "serial_type": "string with leading zero", "cells_and_styles_preserved": cells, "pictures": pictures}


def automatic_workbook(page, url, args, temporary):
    page.goto(url)
    page.wait_for_function("window.ToolApp&&!ToolApp.state.busy")
    page.locator("#workbook").set_input_files(args.workbook)
    page.wait_for_function("ToolApp.state.model&&!ToolApp.state.busy")
    page.evaluate("""text => {
        const parse=ToolOCR.parse;window.torqueRecognizeCalls=0;
        window.ToolOCR={...ToolOCR,recognize:async(file,tools)=>{
            torqueRecognizeCalls++;
            return parse(text,tools,{confidence:95,passes:[{text,confidence:95,source:'original'}]});
        }};
    }""", TEXT)
    photo = temporary / "torque-workflow-fixture.png"
    Image.new("RGB", (400, 250), color=(225, 230, 220)).save(photo)
    page.locator("#photos").set_input_files(photo)
    page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!!ToolApp.state.preparedResult", timeout=30000)
    assert page.evaluate("torqueRecognizeCalls") == 1
    record = page.evaluate("""() => {const p=ToolApp.state.photos[0];return {
        approval:p.approval,serial:p.review.serial,last:p.review.last,due:p.review.due};} """)
    assert record == {"approval": "automatic", "serial": SERIAL, "last": LAST, "due": DUE}, record
    updated = temporary / "torque-calibration-result.xlsx"
    with page.expect_download() as event:
        page.locator("#automatic-result-download").click()
    event.value.save_as(updated)
    return verify_torque_workbook(args.workbook, updated)


def shared_owner_serial_retry(page, url, args, temporary):
    """Run real recognition orchestration, replacing only worker responses.

    A high-confidence owner name cannot resolve two of that person's tools.
    Returning that name first must therefore still lead to a serial crop. The
    previous early-stop decision skipped this and left the torque unmatched.
    """
    page.goto(url)
    page.wait_for_function("window.ToolApp&&!ToolApp.state.busy")
    page.locator("#workbook").set_input_files(args.workbook)
    page.wait_for_function("ToolApp.state.model&&!ToolApp.state.busy")
    page.evaluate("""serial => {
        const tools=ToolApp.state.model.tools,target=tools.find(tool=>tool.serial===serial);
        const ownersTools=tools.filter(tool=>tool.name===target.name);
        if(ownersTools.length<2)throw new Error('Retry fixture requires an owner with multiple tools');
        const dates='교정일자 2026.06.22\\n차기교정예정일자 2026.12.22';
        window.torqueWorkerCalls=[];window.torqueWorkerParameters=[];
        const responses=[target.name+'\\n'+dates,dates,'Serial: '+serial];
        Tesseract.createWorker=async()=>({
            setParameters:async parameters=>{torqueWorkerParameters.push({...parameters});},
            recognize:async canvas=>{
                const index=torqueWorkerCalls.length;
                torqueWorkerCalls.push({width:canvas.width,height:canvas.height});
                return {data:{text:responses[Math.min(index,responses.length-1)],confidence:95,words:[]}};
            },
            terminate:()=>{}
        });
        const input=document.createElement('input');input.type='file';input.id='torque-retry-photo';document.body.append(input);
    }""", SERIAL)
    photo = temporary / "shared-owner-retry-fixture.png"
    Image.new("RGB", (1000, 1400), color="white").save(photo)
    page.locator("#torque-retry-photo").set_input_files(photo)
    result = page.evaluate("""async () => {
        const model=ToolApp.state.model;
        const parsed=await ToolOCR.recognize(document.querySelector('#torque-retry-photo').files[0],model.tools);
        return {assessment:ToolAutomation.assess(model,parsed),failed:parsed.failed,error:parsed.error,
            calls:torqueWorkerCalls,parameters:torqueWorkerParameters,
            passes:parsed.passes.map(pass=>({source:pass.source,text:pass.text})),
            serialMatches:parsed.identityEvidence.serialMatches,
            targetId:model.tools.find(tool=>tool.serial==='0716805829').id};
    }""")
    assert not result["failed"], result
    assert len(result["calls"]) >= 3, result
    assert "identity-crop" in [p["source"] for p in result["passes"]], result
    assert any(match["serial"] == SERIAL and match["confidence"] >= 70 for match in result["serialMatches"]), result
    assessment = result["assessment"]
    assert assessment["eligible"], result
    assert assessment["toolId"] == result["targetId"], result
    assert assessment["serial"] == SERIAL, result
    assert assessment["last"] == LAST and assessment["due"] == DUE, result
    return {"recognize_orchestration_is_real": True, "only_worker_responses_are_controlled": True,
            "shared_owner_does_not_skip_serial_crop": True,
            "worker_recognition_calls": len(result["calls"]), "pass_sources": [p["source"] for p in result["passes"]],
            "matched_serial": assessment["serial"], "Cal Last": assessment["last"], "Cal Next": assessment["due"]}


def main():
    args = arguments(__doc__)
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    errors, remote = [], []
    with tempfile.TemporaryDirectory(prefix="tool-torque-dates-") as directory, serve_program(args.program) as url, sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
        context = browser.new_context(accept_downloads=True, locale="ko-KR")
        context.on("request", lambda request: remote.append(request.url)
                   if urlsplit(request.url).scheme in {"http", "https"} and urlsplit(request.url).netloc != urlsplit(url).netloc else None)
        page = context.new_page()
        page.on("dialog", lambda dialog: dialog.accept())
        page.on("pageerror", lambda error: errors.append(str(error)))
        result = {"torque_label_policy": policy_checks(page, url, args.workbook)}
        result["shared_owner_retry"] = shared_owner_serial_retry(page, url, args, Path(directory))
        result["automatic_excel"] = automatic_workbook(page, url, args, Path(directory))
        assert not errors, errors
        assert not remote, remote
        browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == original_hash
    result.update({"original_company_workbook_unchanged": True, "page_errors": errors, "external_network_requests": remote,
                   "actual_screenshot_ocr_accuracy": "not measured: original JPEG unavailable; OCR text is controlled"})
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
