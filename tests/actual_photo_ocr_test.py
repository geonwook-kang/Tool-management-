"""Run an attached original photograph through real offline OCR and Excel save.

The photograph and workbook are supplied at runtime and remain outside Git.
The original photograph is unchanged by default; optional blur/resolution
variants are marked explicitly in the report. No substitute labels are drawn
and Tesseract responses are never replaced. Expected identity and dates are
explicit test arguments, never inputs to the recognizer. A real browser OPFS
workbook exercises connected-file autosave; actual OneDrive cloud
synchronization is outside this test.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import time
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image, ImageFilter, ImageOps
from playwright.sync_api import sync_playwright

from browser_test import REPO, download_expression, serve_program
from connected_file_test import PICKER_SCRIPT, export_shared, ready, seed_file
from identity_ocr_test import OBSERVE_WORKER
from smax_workbook_test import drawing_info, rel_path


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--photo", type=Path, required=True)
    parser.add_argument("--variant", choices=("original", "mild-blur", "reduced-resolution"), default="original",
                        help="Optional, explicitly generated degradation of the real photo; original bytes are the default")
    parser.add_argument("--serial", required=True)
    parser.add_argument("--name", required=True)
    parser.add_argument("--last", required=True)
    parser.add_argument("--due", required=True)
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--diagnose-only", action="store_true")
    parser.add_argument("--diagnostics", type=Path, default=Path("/tmp/tool-actual-photo-ocr.json"))
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    for label, path in [("program", args.program), ("photo", args.photo), ("workbook", args.workbook)]:
        if path is None or not path.is_file():
            parser.error(f"Supply --{label} with an existing file")
    datetime.fromisoformat(args.last)
    datetime.fromisoformat(args.due)
    return args


def photo_variant(args, temporary):
    if args.variant == "original":
        return args.photo, {"kind": "original", "pixels_or_bytes_altered": False}
    output = temporary / (args.variant + ".jpg")
    with Image.open(args.photo) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        original_dimensions = list(image.size)
        if args.variant == "mild-blur":
            image = image.filter(ImageFilter.GaussianBlur(radius=0.7))
            details = {"gaussian_radius_pixels": 0.7}
        else:
            image.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
            details = {"maximum_edge_pixels": 1800}
        image.save(output, format="JPEG", quality=90)
        dimensions = list(image.size)
    return output, {"kind": args.variant, "pixels_or_bytes_altered": True,
                    "source_dimensions": original_dimensions, "tested_dimensions": dimensions,
                    "jpeg_quality": 90, **details}


def target_reference(workbook, serial):
    book = load_workbook(workbook)
    targets = []
    for sheet in book:
        headers = {cell.value: cell.column for cell in sheet[1]}
        if "Serial" not in headers:
            continue
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headers["Serial"]).value == serial:
                targets.append({"sheet": sheet.title, "row": row, "headers": headers,
                    "last_cell": sheet.cell(row, headers["Cal Last"]).coordinate,
                    "due_cell": sheet.cell(row, headers["Cal Next"]).coordinate,
                    "prior_last": str(sheet.cell(row, headers["Cal Last"]).value),
                    "prior_due": str(sheet.cell(row, headers["Cal Next"]).value)})
    book.close()
    assert len(targets) == 1, "The expected serial must identify exactly one original workbook row"
    return targets[0]


def verify_saved(original, updated, reference, args, prepared_hash):
    before, after = load_workbook(original), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    allowed = {(reference["sheet"], reference["last_cell"]): datetime.fromisoformat(args.last),
               (reference["sheet"], reference["due_cell"]): datetime.fromisoformat(args.due)}
    checked = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for old_row in old_sheet:
            for old in old_row:
                new = new_sheet[old.coordinate]
                expected = allowed.get((old_sheet.title, old.coordinate), old.value)
                assert new.value == expected, (old_sheet.title, old.coordinate, old.value, new.value, expected)
                assert new._style == old._style, (old_sheet.title, old.coordinate)
                checked += 1
    serial_cell = after[reference["sheet"]].cell(reference["row"], reference["headers"]["Serial"])
    assert serial_cell.value == args.serial and serial_cell.data_type == "s"
    pictures = sum(len(sheet._images) for sheet in after)
    assert pictures == sum(len(sheet._images) for sheet in before)
    with ZipFile(original) as old_zip, ZipFile(updated) as new_zip:
        sheet_part, drawing_part, old_pictures = drawing_info(old_zip, reference["sheet"])
        new_sheet_part, new_drawing_part, new_pictures = drawing_info(new_zip, reference["sheet"])
        assert (sheet_part, drawing_part) == (new_sheet_part, new_drawing_part)
        assert set(old_pictures) == set(new_pictures)
        target = (reference["row"], reference["headers"]["Label Photo"])
        for location, old in old_pictures.items():
            new = new_pictures[location]
            if location == target:
                assert len(new) == 1 and sha256(new[0]["bytes"]).hexdigest() == prepared_hash
            else:
                assert old == new, ("Unselected picture or anchor changed", location)
        old_names, new_names = set(old_zip.namelist()), set(new_zip.namelist())
        assert old_names <= new_names
        changed = {name for name in old_names if old_zip.read(name) != new_zip.read(name)}
        allowed_parts = {sheet_part, drawing_part, rel_path(drawing_part)}
        if "[Content_Types].xml" in changed:
            signature = lambda node: (node.tag, tuple(sorted(node.attrib.items())), (node.text or "").strip())
            old_entries = {signature(node) for node in ET.fromstring(old_zip.read("[Content_Types].xml"))}
            new_entries = {signature(node) for node in ET.fromstring(new_zip.read("[Content_Types].xml"))}
            assert old_entries <= new_entries
            assert all(dict(attributes) == {"Extension": "jpg", "ContentType": "image/jpeg"}
                       and tag.endswith("}Default") for tag, attributes, _ in new_entries - old_entries)
            assert len(new_entries - old_entries) <= 1
            allowed_parts.add("[Content_Types].xml")
        assert changed <= allowed_parts, changed - allowed_parts
        added = {name for name in new_names - old_names if not name.endswith("/")}
        assert len(added) == 1 and all(name.startswith("xl/media/") for name in added)
    before.close()
    after.close()
    return {"cells_and_styles_checked": checked, "pictures": pictures,
            "correct_label_photo_anchor_and_bytes": True,
            "other_cells_styles_smax_and_photos_preserved": True,
            "changed_zip_parts": sorted(changed)}


def main():
    args = arguments()
    source_hash = sha256(args.workbook.read_bytes()).hexdigest()
    photo_hash = sha256(args.photo.read_bytes()).hexdigest()
    reference = target_reference(args.workbook, args.serial)
    report = {"program": str(args.program), "program_sha256": sha256(args.program.read_bytes()).hexdigest(),
              "original_photo_sha256": photo_hash, "real_original_photo": args.variant == "original",
              "worker_results_replaced": False, "expected_values_passed_to_recognizer": False,
              "actual_onedrive_sync_tested": False, "reference": reference,
              "diagnose_only": args.diagnose_only, "failures": []}
    with tempfile.TemporaryDirectory(prefix="tool-actual-photo-") as directory, serve_program(args.program) as url, sync_playwright() as playwright:
        temporary = Path(directory)
        tested_photo, variation = photo_variant(args, temporary)
        report["photo_variation"] = variation
        report["tested_photo_sha256"] = sha256(tested_photo.read_bytes()).hexdigest()
        browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
        context = browser.new_context(accept_downloads=True, locale="ko-KR")
        context.add_init_script("(" + PICKER_SCRIPT + ")()")
        page = context.new_page()
        errors, requests = [], []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(url)
        page.wait_for_function("window.ToolApp&&window.ToolOCR&&window.ToolRuntime")
        page.evaluate(OBSERVE_WORKER)
        context.route("**/*", lambda route: (requests.append(route.request.url), route.abort())
                      if route.request.url.startswith(("http:", "https:")) else route.continue_())
        context.set_offline(True)
        baseline_hash = seed_file(page, args.workbook)
        page.locator("#connect-onedrive").click()
        page.wait_for_function("ToolApp.state.source.connected&&ToolApp.state.model")
        ready(page)
        began = time.monotonic()
        page.locator("#photos").set_input_files(tested_photo)
        page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!ToolApp.state.source.saving", timeout=360000)
        report.update(page.evaluate("""async() => {
          const photo=ToolApp.state.photos[0],handle=await(await navigator.storage.getDirectory()).getFileHandle('company-shared.xlsx');
          const backup=await ToolConnectedFile.takeBackup();return {result:photo.result,assessment:photo.assessment,
            approval:photo.approval,saved:!!photo.saved,review:!!photo.review,
            prepared:!!ToolApp.state.preparedResult,preparedPhotoHash:await testHash(photo.image.photoBytes),
            saveCount:ToolApp.state.source.saveCount,sourceHash:await testHash(await(await handle.getFile()).arrayBuffer()),
            backupHash:backup&&await testHash(backup.bytes),sourceError:ToolApp.state.source.error,
            nameInput:document.querySelector('#name').value,serialInput:document.querySelector('#serial').value,
            lastInput:document.querySelector('#last').value,dueInput:document.querySelector('#due').value,
            readings:identityReadings};
        }"""))
        report.update({"seconds": round(time.monotonic() - began, 2), "http_requests": requests, "page_errors": errors})
        result, assessment = report["result"] or {}, report["assessment"] or {}
        if not args.diagnose_only:
            expected = (args.name, args.serial, args.last, args.due)
            # A confidently read unique serial may resolve the registered
            # owner even when the handwritten name remains unreadable.
            if tuple(result.get(key) for key in ("serial", "last", "due")) != (args.serial, args.last, args.due):
                report["failures"].append("OCR did not read the exact serial and both complete photo dates")
            for label, actual in [("assessment", tuple(assessment.get(key) for key in ("name", "serial", "last", "due"))),
                                  ("visible inputs", tuple(report.get(key) for key in ("nameInput", "serialInput", "lastInput", "dueInput")))]:
                if actual != expected:
                    report["failures"].append(f"{label} did not match all four observed photo fields: {actual!r}")
            if result.get("failed") or result.get("identityRecoveryIncomplete"):
                report["failures"].append("Real OCR failed or identity recovery did not finish")
            if not assessment.get("eligible") or not report["saved"] or report["saveCount"] != 1 or report["approval"] != "automatic":
                report["failures"].append("The actual photograph was not automatically saved once")
        if report["saved"]:
            assert assessment.get("photoKind", "calibration") == "calibration"
            assert report["backupHash"] == baseline_hash == source_hash and report["sourceHash"] != baseline_hash
            output = temporary / "saved.xlsx"
            export_shared(page, output)
            report["workbook_verification"] = verify_saved(args.workbook, output, reference, args, report["preparedPhotoHash"])
            backup = temporary / "backup.xlsx"
            download_expression(page, "(await ToolConnectedFile.takeBackup()).bytes", backup)
            assert backup.read_bytes() == args.workbook.read_bytes(), "The original backup must be byte-exact"
            report["original_backup_byte_exact"] = True
        else:
            assert report["sourceHash"] == baseline_hash and report["saveCount"] == 0
        assert not errors and not requests, (errors, requests)
        browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == source_hash
    assert sha256(args.photo.read_bytes()).hexdigest() == photo_hash
    report.update({"original_photo_unchanged": True, "original_company_workbook_unchanged": True})
    args.diagnostics.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("program_sha256", "photo_variation", "seconds", "assessment", "saved", "failures", "http_requests", "page_errors")}
                     | {"workbook": report.get("workbook_verification"), "diagnostics": str(args.diagnostics)}, ensure_ascii=False, indent=2))
    assert not report["failures"], report["failures"]


if __name__ == "__main__":
    main()
