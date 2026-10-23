"""Real offline OCR of the original SMAX screenshots and a calibration photo.

The supplied XLSX is read only. Its actual SMAX screenshots are extracted with
unchanged bytes and processed by real Tesseract; no responses are substituted.
Connected writes use a disposable, real browser OPFS workbook. A temporary
copy starts each target Cal Next one day earlier to demonstrate an actual date
write if its SMAX Due Date is readable. Cal Last and all Label Photos must stay
unchanged. This does not test the user's newly supplied inline-only photograph
or actual OneDrive synchronization.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta
from hashlib import sha256
import json
from pathlib import Path
import re
import tempfile
import time
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from playwright.sync_api import sync_playwright

from browser_test import REPO, serve_program
from connected_file_test import PICKER_SCRIPT, export_shared, ready, seed_file
from identity_ocr_test import OBSERVE_WORKER
from smax_workbook_test import drawing_info, rel_path
from torque_ocr_test import extract_sources


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--diagnose-only", action="store_true")
    parser.add_argument("--diagnostics", type=Path, default=Path("/tmp/tool-smax-real-ocr.json"))
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    if args.workbook is None or not args.workbook.is_file():
        parser.error("Supply --workbook with the original company XLSX.")
    if not args.program.is_file():
        parser.error("Supply --program with the standalone HTML program.")
    return args


def cases_from_original(workbook, temporary):
    cases = []
    book = load_workbook(workbook)
    for serial in ("0716805829", "31550075"):
        directory = temporary / serial
        directory.mkdir()
        reference, sources = extract_sources(SimpleNamespace(workbook=workbook, serial=serial), directory)
        sheet = book[reference["sheet"]]
        headers = {cell.value: cell.column for cell in sheet[1]}
        reference["headers"] = headers
        reference["due_cell"] = sheet.cell(reference["row"], headers["Cal Next"]).coordinate
        with ZipFile(workbook) as archive:
            reference["sheet_part"] = drawing_info(archive, reference["sheet"])[0]
        system = next(item for item in sources if item["kind"] == "system-screenshot")
        anchored = [item for item in sheet._images
                    if item.anchor._from.row == reference["row"] - 1
                    and item.anchor._from.col == headers["SMAX Tool Photo"] - 1]
        assert len(anchored) == 1 and sha256(anchored[0]._data()).hexdigest() == system["sha256"]
        cases.append({**system, "kind": "actual-smax-" + serial, "reference": reference})
        if serial == "0716805829":
            label = next(item for item in sources if item["kind"] == "calibration-label")
            control = {**label, "kind": "actual-calibration-control", "reference": reference}
    book.close()
    return cases + [control]


def earlier_next_fixture(original, output, reference):
    book = load_workbook(original, read_only=True)
    epoch = book.epoch
    book.close()
    prior = datetime.fromisoformat(reference["due"]) - timedelta(days=1)
    numeric = to_excel(prior, epoch)
    value = str(int(numeric)) if numeric.is_integer() else str(numeric)
    with ZipFile(original) as before, ZipFile(output, "w") as after:
        after.comment = before.comment
        for info in before.infolist():
            data = before.read(info.filename)
            if info.filename == reference["sheet_part"]:
                pattern = r'(<c\b[^>]*\br="' + re.escape(reference["due_cell"]) + r'"[^>]*>.*?<v>)([^<]*)(</v>)'
                text, count = re.subn(pattern, lambda match: match[1] + value + match[3], data.decode("utf-8"), flags=re.S)
                assert count == 1
                data = text.encode("utf-8")
            after.writestr(info, data)
    with ZipFile(original) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        assert {name for name in before.namelist() if before.read(name) != after.read(name)} == {reference["sheet_part"]}
    return prior.strftime("%Y-%m-%d")


def verify_saved_workbook(baseline, updated, reference, assessment, prepared_hash):
    before, after = load_workbook(baseline), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    allowed = {(reference["sheet"], reference["due_cell"])} if assessment["updateDue"] else set()
    checked = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for old_row in old_sheet:
            for old in old_row:
                new = new_sheet[old.coordinate]
                expected = datetime.fromisoformat(reference["due"]) if (old_sheet.title, old.coordinate) in allowed else old.value
                assert new.value == expected, (old_sheet.title, old.coordinate, old.value, new.value, expected)
                assert new._style == old._style, (old_sheet.title, old.coordinate)
                checked += 1
    target_sheet = after[reference["sheet"]]
    assert target_sheet.cell(reference["row"], reference["headers"]["Cal Last"]).value.strftime("%Y-%m-%d") == reference["last"]
    serial = target_sheet.cell(reference["row"], reference["headers"]["Serial"])
    assert serial.value == reference["serial"] and serial.data_type == "s"
    pictures = sum(len(sheet._images) for sheet in after)
    assert pictures == sum(len(sheet._images) for sheet in before) == 59
    with ZipFile(baseline) as original_zip, ZipFile(updated) as updated_zip:
        sheet_part, drawing_part, old_pictures = drawing_info(original_zip, reference["sheet"])
        new_sheet_part, new_drawing_part, new_pictures = drawing_info(updated_zip, reference["sheet"])
        assert (sheet_part, drawing_part) == (new_sheet_part, new_drawing_part)
        target = (reference["row"], reference["headers"]["SMAX Tool Photo"])
        assert set(old_pictures) == set(new_pictures)
        for location, old in old_pictures.items():
            new = new_pictures[location]
            if location == target:
                assert len(new) == 1 and sha256(new[0]["bytes"]).hexdigest() == prepared_hash
            else:
                assert old == new, ("Unselected image or anchor changed", location)
        old_names, new_names = set(original_zip.namelist()), set(updated_zip.namelist())
        assert old_names <= new_names
        changed = {name for name in old_names if original_zip.read(name) != updated_zip.read(name)}
        allowed_parts = {drawing_part, rel_path(drawing_part)} | ({sheet_part} if allowed else set())
        if "[Content_Types].xml" in changed:
            # The original workbook registers JPEG as .jpeg; the app's actual
            # image conversion writes .jpg. Permit its required MIME entry,
            # preserving every existing declaration and no unrelated additions.
            old_types = ET.fromstring(original_zip.read("[Content_Types].xml"))
            new_types = ET.fromstring(updated_zip.read("[Content_Types].xml"))
            signature = lambda node: (node.tag, tuple(sorted(node.attrib.items())), (node.text or "").strip())
            old_entries, new_entries = {signature(node) for node in old_types}, {signature(node) for node in new_types}
            assert old_entries <= new_entries
            assert all(dict(attributes) == {"Extension": "jpg", "ContentType": "image/jpeg"}
                       and tag.endswith("}Default") for tag, attributes, _ in new_entries - old_entries)
            assert len(new_entries - old_entries) <= 1
            allowed_parts.add("[Content_Types].xml")
        assert changed <= allowed_parts, changed - allowed_parts
        added = {name for name in new_names - old_names if not name.endswith("/")}
        assert len(added) == 1 and all(name.startswith("xl/media/") for name in added)
        if not assessment["updateDue"]:
            assert original_zip.read(sheet_part) == updated_zip.read(sheet_part)
    before.close()
    after.close()
    return {"cells_and_styles_preserved": checked, "pictures": pictures,
            "cal_last_unchanged": True, "label_photo_and_all_other_photos_unchanged": True,
            "correct_smax_photo_anchor_and_bytes": True, "changed_zip_parts": sorted(changed),
            "due_updated_from_actual_field": assessment["updateDue"]}


def run_case(browser, url, args, temporary, case):
    reference = case["reference"]
    baseline = temporary / (case["kind"] + "-source.xlsx")
    prior = earlier_next_fixture(args.workbook, baseline, reference)
    context = browser.new_context(accept_downloads=True, locale="ko-KR")
    context.add_init_script("(" + PICKER_SCRIPT + ")()")
    page = context.new_page()
    errors, requests = [], []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(url)
    page.wait_for_function("window.ToolApp&&window.ToolSMAX&&window.ToolOCR&&window.ToolRuntime")
    page.evaluate(OBSERVE_WORKER)
    context.route("**/*", lambda route: (requests.append(route.request.url), route.abort())
                  if route.request.url.startswith(("http:", "https:")) else route.continue_())
    context.set_offline(True)
    baseline_hash = seed_file(page, baseline)
    page.locator("#connect-onedrive").click()
    page.wait_for_function("ToolApp.state.source.connected&&ToolApp.state.model")
    ready(page)
    began = time.monotonic()
    page.locator("#photos").set_input_files(case["path"])
    page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!ToolApp.state.source.saving", timeout=360000)
    evidence = page.evaluate("""async() => {
      const photo=ToolApp.state.photos[0],handle=await(await navigator.storage.getDirectory()).getFileHandle('company-shared.xlsx');
      const backup=await ToolConnectedFile.takeBackup();return {result:photo.result,assessment:photo.assessment,
        approval:photo.approval,saved:!!photo.saved,review:!!photo.review,prepared:!!ToolApp.state.preparedResult,
        photoKind:photo.review?.photoKind||photo.assessment?.photoKind,
        preparedPhotoHash:await testHash(photo.image.photoBytes),saveCount:ToolApp.state.source.saveCount,
        sourceHash:await testHash(await(await handle.getFile()).arrayBuffer()),
        backupHash:backup&&await testHash(backup.bytes),sourceError:ToolApp.state.source.error,
        ownerNote:document.querySelector('#photo-kind-note').textContent,
        nameInput:document.querySelector('#name').value,serialInput:document.querySelector('#serial').value,
        readings:identityReadings};
    }""")
    report = {key: value for key, value in case.items() if key != "path"}
    report.update(evidence)
    report.update({"prior_next": prior, "seconds": round(time.monotonic() - began, 2),
                   "http_requests": requests, "page_errors": errors, "failures": []})
    result, assessment = evidence["result"] or {}, evidence["assessment"] or {}
    control = case["kind"] == "actual-calibration-control"
    fields = result.get("smaxEvidence") or {}
    confident_serials = [item for item in fields.get("serialReadings", [])
                         if item.get("labelVerified") and item.get("value") == reference["serial"]
                         and (item.get("confidence") or 0) >= 70]
    correct_dues = [item for item in fields.get("dueReadings", [])
                    if item.get("labelVerified") and item.get("value") == reference["due"]]
    report["field_summary"] = {"correct_serial_confidence": max((item["confidence"] for item in confident_serials), default=None),
        "correct_due_confidence": max((item.get("confidence") or 0 for item in correct_dues), default=None),
        "correct_serial_sources": sorted({item["source"] for item in confident_serials}),
        "correct_due_sources": sorted({item["source"] for item in correct_dues})}
    if not args.diagnose_only:
        if result.get("failed"):
            report["failures"].append("Real Tesseract did not finish")
        if control:
            if result.get("documentType") != "calibration":
                report["failures"].append("The actual calibration label was mistaken for a SMAX screen")
            if (result.get("last"), result.get("due")) != (reference["last"], reference["due"]):
                report["failures"].append("The actual calibration-label date pair regressed")
            assert not evidence["saved"] and not evidence["prepared"]
        else:
            if result.get("documentType") != "smax":
                report["failures"].append("The original SMAX screen was not classified from its field labels")
            if not assessment.get("eligible") or assessment.get("serial") != reference["serial"] or assessment.get("matchMethod") != "smax-serial":
                report["failures"].append("The actual Serial / Asset Number did not identify the registered tool")
            if not evidence["saved"] or evidence["saveCount"] != 1 or evidence["approval"] != "automatic" or evidence["photoKind"] != "smax":
                report["failures"].append("The actual screenshot was not automatically saved to the SMAX slot once")
    if evidence["saved"]:
        assert not control
        assert confident_serials, "The saved tool must have confident, label-bound Serial / Asset Number evidence"
        assert evidence["backupHash"] == baseline_hash and evidence["sourceHash"] != baseline_hash
        assert assessment["name"] == reference["name"] and assessment["serial"] == reference["serial"]
        assert assessment["last"] == reference["last"]
        assert evidence["nameInput"] == reference["name"] and evidence["serialInput"] == reference["serial"]
        owner = assessment.get("ownerEnglish") or ""
        if owner:
            assert owner in evidence["ownerNote"] and reference["name"] in evidence["ownerNote"]
        report["english_owner_displayed_if_observed"] = bool(owner)
        if assessment.get("updateDue"):
            assert assessment["due"] == reference["due"], ("Wrong SMAX Due Date", assessment, reference)
            assert any((item.get("confidence") or 0) >= 70 for item in correct_dues), "Cal Next needs confident, label-bound Due Date evidence"
        else:
            assert assessment["due"] == prior
        output = temporary / (case["kind"] + "-saved.xlsx")
        export_shared(page, output)
        report["workbook_verification"] = verify_saved_workbook(baseline, output, reference, assessment, evidence["preparedPhotoHash"])
    else:
        assert evidence["sourceHash"] == baseline_hash and evidence["saveCount"] == 0
    assert not errors and not requests, (errors, requests)
    context.close()
    return report


def main():
    args = arguments()
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    diagnostics = {"program": str(args.program), "program_sha256": sha256(args.program.read_bytes()).hexdigest(),
                   "original_image_bytes_preserved": True, "recognizer_responses_replaced": False,
                   "new_inline_user_photo_tested": False, "actual_onedrive_sync_tested": False,
                   "diagnose_only": args.diagnose_only, "cases": []}
    with tempfile.TemporaryDirectory(prefix="tool-smax-real-ocr-") as directory, serve_program(args.program) as url, sync_playwright() as playwright:
        temporary = Path(directory)
        cases = cases_from_original(args.workbook, temporary)
        browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
        for case in cases:
            report = run_case(browser, url, args, temporary, case)
            diagnostics["cases"].append(report)
            args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
            print(json.dumps({"kind": report["kind"], "media": report["media"],
                              "document_type": report["result"].get("documentType"),
                              "assessment": report["assessment"], "saved": report["saved"],
                              "field_evidence": report["field_summary"],
                              "seconds": report["seconds"], "workbook": report.get("workbook_verification"),
                              "failures": report["failures"]}, ensure_ascii=False, indent=2), flush=True)
        browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == original_hash
    failures = [case["kind"] + ": " + failure for case in diagnostics["cases"] for failure in case["failures"]]
    diagnostics.update({"failures": failures, "original_company_workbook_unchanged": True})
    args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"real_smax_original_cases": 2, "real_calibration_control_cases": 1,
                      "new_inline_user_photo_tested": False, "failures": failures,
                      "original_company_workbook_unchanged": True, "diagnostics": str(args.diagnostics)}, ensure_ascii=False, indent=2))
    assert not failures, failures


if __name__ == "__main__":
    main()
