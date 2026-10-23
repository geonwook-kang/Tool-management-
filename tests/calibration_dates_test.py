"""Verify photographed calibration labels reach the matching Excel date headers.

The examples describe observed label layouts, including Korean table headings
split over lines and bilingual/date noise. Recognition supplies controlled OCR
text to the real parser/decision code. The connected case writes a real browser
OPFS workbook, then verifies Cal Last and Cal Next by their column headings.
Actual camera-photo OCR accuracy is covered separately by the real OCR tests.
The uploaded company workbook is read only; outputs remain temporary.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import tempfile
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright

from browser_test import REPO, arguments, download_expression, fixture_photo, serve_program, verify_workbook
from connected_file_test import PICKER_SCRIPT, export_shared, ready, seed_file


LAST, DUE = "2025-12-29", "2026-12-29"
IDENTITY = "강건욱\n31550075\n"


def observed_cases():
    labelled = "교정일자 2025.12.29\n차기교정예정일자 2026.12.29"
    split = "교정\n일자\n2025.12.29\n차기교정\n예정일자\n2026.12.29"
    table_text = IDENTITY + "교정일자\n차기교정예정일자\n2025.12.29\n2026.12.29"
    table_words = [
        {"text": "강건욱", "confidence": 95, "bbox": {"x0": 20, "y0": 10, "x1": 180, "y1": 45}},
        {"text": "31550075", "confidence": 95, "bbox": {"x0": 20, "y0": 55, "x1": 220, "y1": 90}},
        {"text": "교정일자", "confidence": 95, "bbox": {"x0": 20, "y0": 110, "x1": 240, "y1": 145}},
        {"text": "차기교정예정일자", "confidence": 95, "bbox": {"x0": 20, "y0": 180, "x1": 240, "y1": 215}},
        {"text": "2025.12.29", "confidence": 95, "bbox": {"x0": 300, "y0": 110, "x1": 540, "y1": 145}},
        {"text": "2026.12.29", "confidence": 95, "bbox": {"x0": 300, "y0": 180, "x1": 540, "y1": 215}},
    ]
    noisy_serial_text = IDENTITY + labelled + "\n69927"
    noisy_serial_words = [{"text": text, "confidence": confidence} for text, confidence in [
        ("강건욱", 95), ("31550075", 95), ("교정일자", 95), ("2025.12.29", 95),
        ("차기교정예정일자", 95), ("2026.12.29", 95), ("69927", 15.1),
    ]]
    cases = {
        "korean_label_date_same_line": {"text": IDENTITY + labelled, "eligible": True},
        "korean_table_headings_split_over_lines": {"text": IDENTITY + split, "eligible": True},
        "spaced_korean_bilingual_table": {"text": IDENTITY + "교 정 일 자\nDate of Calibration\n2025.12.29\n차 기 교 정 예 정 일 자\nThe due date of\nnext Calibration\n2026.12.29", "eligible": True},
        "due_field_before_last_field": {"text": IDENTITY + "차기교정예정일자 2026.12.29\n교정일자 2025.12.29", "eligible": True},
        "certificate_and_manufacture_date_noise": {"text": IDENTITY + "성적서번호 DA25-1041-007\n제조일 2024.01.01\n" + labelled, "eligible": True},
        "standalone_label_vocabulary_noise": {"text": "31550075\n예정일\n인정기관\n" + labelled, "eligible": True},
        "standalone_handwritten_name_variant_with_exact_serial": {"text": "장건욱\n31550075\n" + labelled, "eligible": True},
        "column_major_ocr_text_with_same_row_word_positions": {"text": table_text, "passes": [
            {"text": table_text, "confidence": 95, "words": table_words}], "eligible": True},
        "column_major_text_reversed_word_array_positions": {"text": table_text, "passes": [
            {"text": table_text, "confidence": 95, "words": list(reversed(table_words))}], "eligible": True},
        "explicit_due_overrides_workbook_twelve_month_cycle": {"text": IDENTITY + "교정일자 2026.01.15\n차기교정예정일자 2026.07.15", "eligible": True,
                                                               "last": "2026-01-15", "due": "2026-07-15"},
        "unreadable_last_with_visible_label": {"text": IDENTITY + "교정일자 [읽을 수 없음]\n차기교정예정일자 2026.12.29", "eligible": False},
        "unreadable_due_with_visible_label": {"text": IDENTITY + "교정일자 2025.12.29\n차기교정예정일자 [읽을 수 없음]", "eligible": False},
        "invalid_due_month": {"text": IDENTITY + "교정일자 2025.12.29\n차기교정예정일자 2026.13.29", "eligible": False},
        "invalid_last_day": {"text": IDENTITY + "교정일자 2025.02.30\n차기교정예정일자 2026.02.28", "eligible": False},
        "due_earlier_than_last": {"text": IDENTITY + "교정일자 2026.12.29\n차기교정예정일자 2025.12.29", "eligible": False},
        "low_confidence_dates": {"text": IDENTITY + labelled, "confidence": 20, "eligible": False},
        "different_due_dates_between_reads": {"text": IDENTITY + labelled, "passes": [
            {"text": IDENTITY + labelled, "confidence": 95},
            {"text": IDENTITY + "교정일자 2025.12.29\n차기교정예정일자 2026.12.28", "confidence": 95},
        ], "eligible": False},
        "crop_evidence_stays_separate_when_reparsed": {"text": IDENTITY + labelled, "passes": [
            {"text": IDENTITY + labelled, "confidence": 95, "source": "original"},
            {"text": "이름: __OTHER_ENGINEER__\n20281220\n" + labelled, "confidence": 95, "source": "date-crop"},
            {"text": IDENTITY + "교정일자 2028.12.20\n차기교정예정일자 2029.12.20", "confidence": 95, "source": "identity-crop"},
            # An independent repeated crop can produce exactly the same text.
            # The displayed aggregate deduplicates it; replay must retain each
            # pass's purpose instead of treating aggregate text as a new pass.
            {"text": "이름: __OTHER_ENGINEER__\n20281220\n" + labelled, "confidence": 95, "source": "date-crop-repeat"},
        ], "eligible": True, "replay": True},
        "low_confidence_numeric_noise_with_clear_serial": {"text": noisy_serial_text, "passes": [
            {"text": noisy_serial_text, "confidence": 95, "source": "original", "words": noisy_serial_words}], "eligible": True},
        "compact_recognized_due_is_date_not_unknown_serial": {"text": IDENTITY + labelled + "\n20261229", "eligible": True},
        "high_confidence_other_explicit_serial": {"text": IDENTITY + "시리얼 번호: 69927\n" + labelled, "eligible": False},
        "explicit_serial_matching_compact_due_is_still_identity": {"text": IDENTITY + "S/N: 20261229\n" + labelled, "eligible": False},
        "known_other_engineer_with_exact_serial": {"text": "__OTHER_ENGINEER__\n31550075\n" + labelled, "eligible": False},
        "explicit_other_name_with_exact_serial": {"text": "이름: 김민수\n31550075\n" + labelled, "eligible": False},
    }
    return cases


def parser_policy(page, url, workbook):
    """Use current source modules so policy checks do not depend on bundling."""
    address = urlsplit(url)
    page.route("**/calibration-date-policy.html", lambda route: route.fulfill(
        status=200, content_type="text/html", body='<!doctype html><input id="fixture" type="file">'))
    page.goto(f"{address.scheme}://{address.netloc}/calibration-date-policy.html")
    for source in ["vendor/jszip.min.js", "workbook.js", "ocr.js", "automation.js"]:
        page.add_script_tag(path=REPO / "web" / source)
    page.locator("#fixture").set_input_files(workbook)
    outcomes = page.evaluate("""async cases => {
        const model=await ToolWorkbook.load(await document.querySelector('#fixture').files[0].arrayBuffer());
        const target=model.tools.find(tool=>tool.sheet==='광주'&&String(tool.serial)==='31550075');
        const other=model.tools.find(tool=>tool.name&&tool.name!==target.name);
        if(!other)throw new Error('The company fixture needs a second engineer for conflict coverage.');
        const fixtureText=text=>text.replaceAll('__OTHER_ENGINEER__',other.name);
        const output={};
        for(const [name,spec] of Object.entries(cases)) {
            const confidence=spec.confidence??95;
            const text=fixtureText(spec.text);
            const passes=spec.passes?spec.passes.map(pass=>({...pass,text:fixtureText(pass.text)})):
                [{text,confidence,source:'whole-label'}];
            const result=ToolOCR.parse(text,model.tools,{confidence,passes});
            const assessment=ToolAutomation.assess(model,result);
            output[name]={...assessment,expectedTool:target.id,dateEvidence:result.dateEvidence};
            if(spec.replay){
                const replay=ToolOCR.parse(result.text,model.tools,{confidence:result.confidence,passes:result.passes});
                output[name].identityEvidence=result.identityEvidence;
                output[name].replay={assessment:ToolAutomation.assess(model,replay),
                    identityEvidence:replay.identityEvidence,dateEvidence:replay.dateEvidence};
            }
        }
        return output;
    }""", observed_cases())
    for name, spec in observed_cases().items():
        outcome = outcomes[name]
        assert outcome["eligible"] is spec["eligible"], (name, outcome)
        if spec["eligible"]:
            assert outcome["toolId"] == outcome["expectedTool"], (name, outcome)
            assert outcome["last"] == spec.get("last", LAST), (name, outcome)
            assert outcome["due"] == spec.get("due", DUE), (name, outcome)
            assert outcome["dateMethod"] == "labelled", (name, outcome)
        else:
            assert outcome["reasons"], (name, outcome)
        if spec.get("replay"):
            replay = outcome["replay"]
            assert replay["assessment"]["eligible"] is True, (name, replay)
            for field in ["toolId", "name", "serial", "last", "due"]:
                assert replay["assessment"][field] == outcome[field], (name, field, outcome)
            assert replay["identityEvidence"] == outcome["identityEvidence"], (name, outcome)
            assert replay["dateEvidence"] == outcome["dateEvidence"], (name, outcome)
            assert outcome["identityEvidence"]["observedSerials"] == ["31550075"], (name, outcome)
            assert outcome["dateEvidence"]["dates"] == [LAST, DUE], (name, outcome)
    return {"accepted_label_layouts": [name for name, spec in observed_cases().items() if spec["eligible"]],
            "unsafe_or_incomplete_labels_pending": [name for name, spec in observed_cases().items() if not spec["eligible"]]}


def install_label_results(page, specs):
    page.evaluate("""specs => {
        const parse=ToolOCR.parse;
        window.calibrationLabelCalls=[];
        window.ToolOCR={...ToolOCR,recognize:async(file,tools)=>{
            calibrationLabelCalls.push(file.name);
            const spec=specs[file.name];
            if(!spec)throw new Error('Missing label fixture: '+file.name);
            const confidence=spec.confidence??95;
            return parse(spec.text,tools,{confidence,passes:spec.passes||[{text:spec.text,confidence,source:'whole-label'}]});
        }};
    }""", specs)


def assert_headers(path, last, due):
    workbook = load_workbook(path)
    sheet = workbook["광주"]
    headings = {cell.value: cell.column for cell in sheet[1]}
    row = next(row for row in range(2, sheet.max_row + 1)
               if str(sheet.cell(row, headings["Serial"]).value) == "31550075")
    assert sheet.cell(row, headings["Cal Last"]).value == datetime.fromisoformat(last)
    assert sheet.cell(row, headings["Cal Next"]).value == datetime.fromisoformat(due)
    workbook.close()
    return {"engineer_row": row, "Cal Last": last, "Cal Next": due}


def shuffled_header_mapping(page, args, temporary):
    """An alternate company layout cannot accidentally pass fixed A/B writes."""
    fixture = temporary / "shuffled-calibration-headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "광주"
    sheet.append(["Serial", "Cal Next", "Cal Cycle", "Asset #", "Cal Interval", "Label Photo", "Cal Last"])
    sheet.append(["31550075", datetime(2026, 1, 15), "MONTHS", "South_광주_CT_강건욱", 12, None, datetime(2025, 1, 15)])
    sheet["B2"].number_format = sheet["G2"].number_format = "yyyy-mm-dd"
    workbook.create_sheet("Notes")["A1"] = "Unrelated sheet preserved"
    workbook.save(fixture)
    workbook.close()
    before_hash = sha256(fixture.read_bytes()).hexdigest()
    photo = fixture_photo(args, temporary)
    with Image.open(photo) as image:
        width, height = image.size
    page.locator("#fixture").set_input_files(fixture)
    page.evaluate("() => {const input=document.createElement('input');input.type='file';input.id='date-photo';document.body.append(input);}")
    page.locator("#date-photo").set_input_files(photo)
    assessment = page.evaluate("""async ({width,height}) => {
        const model=await ToolWorkbook.load(await document.querySelector('#fixture').files[0].arrayBuffer());
        const text='강건욱\\n31550075\\n교정일자 2026.01.15\\n차기교정예정일자 2026.07.15';
        const parsed=ToolOCR.parse(text,model.tools,{confidence:95,passes:[{text,confidence:95}]});
        const assessment=ToolAutomation.assess(model,parsed);
        window.shuffledModel=model;
        window.shuffledUpdate={toolId:assessment.toolId,name:assessment.name,serial:assessment.serial,
            last:assessment.last,due:assessment.due,photoBytes:new Uint8Array(await document.querySelector('#date-photo').files[0].arrayBuffer()),
            photoExtension:'png',photoWidth:width,photoHeight:height};
        return {assessment,cycle:model.tools[0].months,validation:ToolWorkbook.validate(model,shuffledUpdate)};
    }""", {"width": width, "height": height})
    assert assessment["cycle"] == 12, assessment
    assert assessment["assessment"]["eligible"], assessment
    assert assessment["assessment"]["dateMethod"] == "labelled", assessment
    assert assessment["validation"] == [], assessment
    updated = temporary / "shuffled-calibration-result.xlsx"
    download_expression(page, "await ToolWorkbook.apply(shuffledModel,[shuffledUpdate])", updated)
    headers = assert_headers(updated, "2026-01-15", "2026-07-15")
    after = load_workbook(updated)
    assert after["광주"]["A2"].value == "31550075"
    assert after["광주"]["E2"].value == 12
    assert len(after["광주"]._images) == 1
    assert after["Notes"]["A1"].value == "Unrelated sheet preserved"
    after.close()
    assert sha256(fixture.read_bytes()).hexdigest() == before_hash
    return {"Cal Next_column": "B", "Cal Last_column": "G", "saved_dates_by_header": headers,
            "explicit_six_month_due_preserved_with_twelve_month_cycle": True}


def connected_date_mapping(context, page, url, args, temporary):
    context.add_init_script("(" + PICKER_SCRIPT + ")()")
    page.goto(url)
    page.wait_for_function("window.ToolApp&&!ToolApp.state.busy")
    seed_file(page, args.workbook)
    page.locator("#connect-onedrive").click()
    page.wait_for_function("ToolApp.state.source.connected&&!!ToolApp.state.model")
    ready(page)
    original = temporary / "connected-before.xlsx"
    export_shared(page, original)
    assert original.read_bytes() == args.workbook.read_bytes()

    photo = fixture_photo(args, temporary)
    text = "강건욱\n31550075\n성적서번호 DA26-1041-007\n차기교정\n예정일자\n2027.12.29\n교정\n일자\n2026.12.29"
    install_label_results(page, {photo.name: {"text": text}})
    assert page.locator("#automatic-mode").is_checked()
    page.locator("#photos").set_input_files(photo)
    page.wait_for_function("ToolApp.state.photos.length===1&&ToolApp.state.photos[0].saved&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!ToolApp.state.source.saving", timeout=30000)
    assert page.evaluate("calibrationLabelCalls") == [photo.name]
    record = page.evaluate("""() => {const p=ToolApp.state.photos[0];return {
        approval:p.approval,last:p.review.last,due:p.review.due,dateMethod:p.assessment.dateMethod,
        sourceError:ToolApp.state.source.error,pickerCalls:testPickerCalls};} """)
    assert record == {"approval": "automatic", "last": "2026-12-29", "due": "2027-12-29",
                      "dateMethod": "labelled", "sourceError": "", "pickerCalls": 1}, record
    updated = temporary / "connected-calibration-dates.xlsx"
    export_shared(page, updated)
    headers = assert_headers(updated, "2026-12-29", "2027-12-29")
    integrity = verify_workbook(args.workbook, updated)

    # A later photo with an unreadable next-calibration field must not replace
    # the correctly saved dates using an interval guess or a different date.
    invalid_photo = temporary / "unreadable-next-calibration.png"
    invalid_photo.write_bytes(photo.read_bytes())
    invalid_text = "31550075\n교정일자 2027.12.29\n차기교정예정일자 2028.13.29"
    install_label_results(page, {invalid_photo.name: {"text": invalid_text}})
    page.locator("#photos").set_input_files(invalid_photo)
    page.wait_for_function("ToolApp.state.photos.length===2&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!ToolApp.state.source.saving", timeout=30000)
    assert page.evaluate("!ToolApp.state.photos[1].review&&!ToolApp.state.photos[1].saved")
    assert page.evaluate("ToolApp.state.photos[1].assessment.reasons.length>0")
    after_invalid = temporary / "connected-after-unreadable-label.xlsx"
    export_shared(page, after_invalid)
    assert after_invalid.read_bytes() == updated.read_bytes()
    assert_headers(after_invalid, "2026-12-29", "2027-12-29")
    return {"upload_without_date_entry_or_approval": True, "connected_opfs_headers": headers,
            "saved_workbook_integrity": integrity, "invalid_followup_leaves_file_byte_identical": True}


def main():
    args = arguments(__doc__)
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    errors, remote = [], []
    with tempfile.TemporaryDirectory(prefix="tool-calibration-date-") as directory, serve_program(args.program) as url, sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
        context = browser.new_context(accept_downloads=True, locale="ko-KR")
        context.on("request", lambda request: remote.append(request.url)
                   if urlsplit(request.url).scheme in {"http", "https"} and urlsplit(request.url).netloc != urlsplit(url).netloc else None)
        page = context.new_page()
        page.on("dialog", lambda dialog: dialog.accept())
        page.on("pageerror", lambda error: errors.append(str(error)))
        result = {"label_date_policy": parser_policy(page, url, args.workbook)}
        result["shuffled_excel_headers"] = shuffled_header_mapping(page, args, Path(directory))
        result["connected_excel"] = connected_date_mapping(context, page, url, args, Path(directory))
        assert not errors, errors
        assert not remote, remote
        browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == original_hash
    result.update({"original_workbook_unchanged": True, "page_errors": errors, "external_network_requests": remote,
                   "actual_photo_ocr": "tested separately; OCR text is controlled in these mapping tests"})
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
