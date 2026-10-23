"""Real offline OCR and OPFS save for an explicitly synthetic partial tag.

Only seven actual digits of the workbook serial are printed, without an owner
name. Real Tesseract must independently read that literal fragment twice; the
production resolver then selects the unique workbook serial and its six-month
cycle. This generated clear label is not the user's new blurred photograph.
The picker supplies a real browser OPFS file, not a live OneDrive workbook.
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import time

from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageStat
from playwright.sync_api import sync_playwright

from browser_test import REPO, serve_program
from connected_file_test import PICKER_SCRIPT, export_shared, ready, seed_file
from identity_ocr_test import FONT, OBSERVE_WORKER
from torque_dates_test import LAST, DUE, verify_torque_workbook


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--serial", default="0716805829")
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--diagnose-only", action="store_true")
    parser.add_argument("--diagnostics", type=Path, default=Path("/tmp/tool-partial-identity-ocr.json"))
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    if args.workbook is None or not args.workbook.is_file():
        parser.error("Supply --workbook with the attached company XLSX.")
    if not args.program.is_file():
        parser.error("Supply --program with the standalone HTML program.")
    if not FONT.is_file():
        parser.error("Install fonts-noto-cjk to generate this printed fixture.")
    return args


def reference(workbook, serial):
    book = load_workbook(workbook)
    targets = []
    for sheet in book:
        headers = {str(cell.value): cell.column for cell in sheet[1]}
        if "Serial" not in headers:
            continue
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, headers["Serial"]).value) == serial:
                targets.append({"sheet": sheet.title, "row": row, "serial": serial,
                                "name": str(sheet.cell(row, headers["Asset #"]).value).rsplit("_", 1)[-1].strip(),
                                "photo_column": headers["Label Photo"] - 1})
    book.close()
    assert len(targets) == 1, "The source workbook must contain one unique target serial"
    return targets[0]


def synthetic_label(output, fragment):
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    numbers = ImageFont.truetype(str(FONT), 94)
    labels = ImageFont.truetype(str(FONT), 58)
    # No canonical serial, owner name, certificate number, or other identity
    # is printed. The resolver must retain the distinction between these seven
    # observed digits and the full serial supplied by the reference workbook.
    draw.text((95, 70), fragment, font=numbers, fill="black")
    draw.text((95, 300), "Date of Calibration", font=labels, fill="black")
    draw.text((95, 405), LAST.replace("-", "."), font=numbers, fill="black")
    draw.text((95, 640), "Due date of next Calibration", font=labels, fill="black")
    draw.text((95, 745), DUE.replace("-", "."), font=numbers, fill="black")
    image.save(output)


def main():
    args = arguments()
    initial_hash = sha256(args.workbook.read_bytes()).hexdigest()
    target = reference(args.workbook, args.serial)
    fragment = args.serial[-7:]
    assert len(fragment) >= 6 and len(fragment) / len(args.serial) >= 0.6
    diagnostics = {"program": str(args.program), "program_sha256": sha256(args.program.read_bytes()).hexdigest(),
                   "synthetic_printed_label": True, "new_user_photo_tested": False,
                   "worker_results_replaced": False, "actual_onedrive_sync_tested": False,
                   "printed_fragment": fragment, "workbook_reference": target}
    errors, requests, failures = [], [], []
    with tempfile.TemporaryDirectory(prefix="tool-partial-real-ocr-") as directory:
        temporary = Path(directory)
        photo = temporary / "printed-partial-serial.png"
        synthetic_label(photo, fragment)
        diagnostics["photo_sha256"] = sha256(photo.read_bytes()).hexdigest()
        with serve_program(args.program) as url, sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
            context = browser.new_context(accept_downloads=True, locale="ko-KR")
            context.add_init_script("(" + PICKER_SCRIPT + ")()")
            page = context.new_page()
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(url)
            page.wait_for_function("window.ToolApp&&window.ToolOCR&&window.ToolRuntime")
            page.evaluate(OBSERVE_WORKER)
            context.route("**/*", lambda route: (requests.append(route.request.url), route.abort())
                          if route.request.url.startswith(("http:", "https:")) else route.continue_())
            context.set_offline(True)
            baseline = seed_file(page, args.workbook)
            assert baseline == initial_hash
            page.locator("#connect-onedrive").click()
            page.wait_for_function("ToolApp.state.source.connected&&ToolApp.state.model")
            ready(page)
            assert page.locator("#automatic-mode").is_checked()
            began = time.monotonic()
            page.locator("#photos").set_input_files(photo)
            page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing&&!ToolApp.state.source.saving", timeout=360000)
            evidence = page.evaluate("""async() => {
              const photo=ToolApp.state.photos[0],handle=await(await navigator.storage.getDirectory()).getFileHandle('company-shared.xlsx');
              const backup=await ToolConnectedFile.takeBackup();
              return {result:photo.result,assessment:photo.assessment,approval:photo.approval,
                review:photo.review&&{toolId:photo.review.toolId,serial:photo.review.serial,
                  last:photo.review.last,due:photo.review.due},
                image:{sha256:await testHash(photo.image.photoBytes),extension:photo.image.photoExtension,
                  width:photo.image.photoWidth,height:photo.image.photoHeight},
                saved:photo.saved,prepared:!!ToolApp.state.preparedResult,
                saveCount:ToolApp.state.source.saveCount,sourceError:ToolApp.state.source.error,
                sourceHash:await testHash(await(await handle.getFile()).arrayBuffer()),
                backupHash:backup&&await testHash(backup.bytes),readings:identityReadings};
            }""")
            diagnostics.update(evidence)
            diagnostics["seconds"] = round(time.monotonic() - began, 2)
            args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
            result, assessment = evidence["result"] or {}, evidence["assessment"] or {}
            identity = result.get("identityEvidence") or {}
            high_sources = set()
            for proof in identity.get("observedSerialEvidence", []):
                if proof.get("serial") != fragment or (proof.get("confidence") or 0) < 70:
                    continue
                for reading in proof.get("sourceReadings", []):
                    source = reading.get("source", "")
                    if (reading.get("confidence") or 0) >= 70 and source in proof.get("passSources", []) and not source.startswith(("date-", "label-")):
                        high_sources.add(source)
            diagnostics["actual_confident_fragment_sources"] = sorted(high_sources)
            # The full leading-zero serial must come from Excel after resolving
            # the partial tag. It must never be claimed as observed OCR text.
            assert not any(args.serial in item.get("text", "") for item in result.get("passes", [])), result
            assert not identity.get("serialMatches"), identity
            assert not identity.get("nameMatches"), identity
            if not args.diagnose_only:
                if result.get("failed"):
                    failures.append("Real OCR did not complete")
                if len(high_sources) < 2:
                    failures.append("The literal fragment was not independently read twice at confidence >=70")
                if not assessment.get("eligible") or assessment.get("matchMethod") != "partial-serial":
                    failures.append("Production resolver did not approve the actual partial evidence")
                if not evidence["saved"] or evidence["saveCount"] != 1 or not evidence["prepared"]:
                    failures.append("The connected OPFS workbook was not saved exactly once")
                if evidence["approval"] != "automatic":
                    failures.append("The generated label required manual approval")
                if evidence["saved"]:
                    assert assessment["serial"] == args.serial and assessment["name"] == target["name"]
                    assert (assessment["last"], assessment["due"]) == (LAST, DUE)
                    assert evidence["backupHash"] == initial_hash and evidence["sourceHash"] != initial_hash
                    updated = temporary / "partial-serial-connected-result.xlsx"
                    export_shared(page, updated)
                    diagnostics["workbook_verification"] = verify_torque_workbook(args.workbook, updated)
                    saved_book = load_workbook(updated)
                    anchored = [item for item in saved_book[target["sheet"]]._images
                                if item.anchor._from.row == target["row"] - 1 and item.anchor._from.col == target["photo_column"]]
                    assert len(anchored) == 1, "The target must have one photo at its existing position"
                    saved_photo = anchored[0]._data()
                    assert sha256(saved_photo).hexdigest() == evidence["image"]["sha256"], "The actual app-prepared photo bytes must be saved at the target"
                    # The app intentionally compresses uploads to JPEG. Verify
                    # that its stored image still depicts the supplied fixture.
                    with Image.open(BytesIO(saved_photo)) as encoded, Image.open(photo) as original:
                        assert encoded.size == original.size == (1800, 1050)
                        difference = ImageStat.Stat(ImageChops.difference(encoded.convert("RGB"), original.convert("RGB"))).mean
                        assert max(difference) < 3, difference
                    diagnostics["saved_photo_verification"] = {"correct_anchor": True,
                        "matches_app_prepared_bytes": True, "production_format": evidence["image"]["extension"],
                        "mean_pixel_difference_after_jpeg": difference}
                    saved_book.close()
            if not evidence["saved"]:
                assert evidence["sourceHash"] == initial_hash and evidence["saveCount"] == 0
            assert not requests, requests
            assert not errors, errors
            browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == initial_hash, "Original company workbook changed"
    diagnostics.update({"failures": failures, "http_requests": requests, "page_errors": errors, "original_unchanged": True})
    args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"synthetic_printed_label": True, "new_user_photo_tested": False,
                      "actual_confident_fragment_sources": sorted(high_sources),
                      "assessment": assessment, "connected_opfs_save": bool(evidence["saved"]),
                      "seconds": diagnostics["seconds"], "workbook": diagnostics.get("workbook_verification"),
                      "failures": failures, "http_requests": requests, "original_unchanged": True,
                      "diagnostics": str(args.diagnostics)}, ensure_ascii=False, indent=2))
    assert not failures, failures


if __name__ == "__main__":
    main()
