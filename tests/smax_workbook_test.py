"""SMAX and calibration photo slots: real XLSX and Chromium OPFS regression tests.

These tests exercise the writer and real browser storage directly. OCR and
classification are outside this suite. The attached company workbook is read
only; photos, reordered-column fixtures and output workbooks stay temporary.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import posixpath
import tempfile
from urllib.parse import urlsplit
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorkbookImage
from PIL import Image, ImageDraw
from playwright.sync_api import sync_playwright

from browser_test import arguments, serve_program
from connected_file_test import PICKER_SCRIPT, export_shared, seed_file
from torque_dates_test import SERIAL, LAST, DUE


NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/package/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def xml_signature(node):
    return (node.tag, tuple(sorted(node.attrib.items())), (node.text or "").strip(),
            tuple(xml_signature(child) for child in node))


def resolve_part(base, target):
    return posixpath.normpath(target.lstrip("/") if target.startswith("/")
                             else posixpath.join(posixpath.dirname(base), target))


def rel_path(part):
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def rels(archive, part):
    return {node.attrib["Id"]: resolve_part(part, node.attrib["Target"])
            for node in ET.fromstring(archive.read(rel_path(part)))
            if node.attrib.get("TargetMode") != "External"}


def drawing_info(archive, sheet_name):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    sheets = workbook.find("s:sheets", NS)
    sheet = next(node for node in sheets if node.attrib["name"] == sheet_name)
    sheet_part = rels(archive, "xl/workbook.xml")[sheet.attrib[f"{{{NS['r']}}}id"]]
    sheet_xml = ET.fromstring(archive.read(sheet_part))
    drawing = sheet_xml.find("s:drawing", NS)
    drawing_part = rels(archive, sheet_part)[drawing.attrib[f"{{{NS['r']}}}id"]]
    drawing_rels = rels(archive, drawing_part)
    pictures = {}
    for anchor in ET.fromstring(archive.read(drawing_part)):
        origin = anchor.find("xdr:from", NS)
        if origin is None or anchor.find("xdr:pic", NS) is None:
            continue
        location = (int(origin.find("xdr:row", NS).text) + 1,
                    int(origin.find("xdr:col", NS).text) + 1)
        embed = anchor.find(".//a:blip", NS).attrib[f"{{{NS['r']}}}embed"]
        media = drawing_rels[embed]
        pictures.setdefault(location, []).append({"anchor": xml_signature(anchor),
            "media": media, "bytes": archive.read(media)})
    return sheet_part, drawing_part, pictures


def verify_output(original, updated, smax_photo, calibration_photo=None, due=None, last=None):
    before, after = load_workbook(original), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    target = None
    for sheet in before:
        headings = {cell.value: cell.column for cell in sheet[1]}
        if "Serial" not in headings:
            continue
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headings["Serial"]).value == SERIAL:
                assert target is None
                target = (sheet.title, row, headings)
    assert target is not None
    title, row, headings = target
    allowed = {}
    for heading, value in [("Cal Next", due), ("Cal Last", last)]:
        if value is not None:
            allowed[(title, after[title].cell(row, headings[heading]).coordinate)] = datetime.fromisoformat(value)
    checked = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for old_row in old_sheet:
            for old in old_row:
                new = new_sheet[old.coordinate]
                key = (old_sheet.title, old.coordinate)
                assert new.value == allowed.get(key, old.value), key
                assert new._style == old._style, key
                checked += 1
    assert after[title].cell(row, headings["Serial"]).data_type == "s"
    assert sum(len(sheet._images) for sheet in before) == sum(len(sheet._images) for sheet in after)
    with ZipFile(original) as old_zip, ZipFile(updated) as new_zip:
        sheet_part, drawing_part, old_pictures = drawing_info(old_zip, title)
        new_sheet_part, new_drawing_part, new_pictures = drawing_info(new_zip, title)
        assert (sheet_part, drawing_part) == (new_sheet_part, new_drawing_part)
        changed_slots = {}
        if smax_photo is not None:
            changed_slots[(row, headings["SMAX Tool Photo"])] = smax_photo.read_bytes()
        if calibration_photo is not None:
            changed_slots[(row, headings["Label Photo"])] = calibration_photo.read_bytes()
        assert set(old_pictures) == set(new_pictures)
        for location, old in old_pictures.items():
            new = new_pictures[location]
            if location in changed_slots:
                assert len(new) == 1 and new[0]["bytes"] == changed_slots[location], location
            else:
                assert old == new, ("Unselected photo or anchor changed", location)
        old_names, new_names = set(old_zip.namelist()), set(new_zip.namelist())
        assert old_names <= new_names
        changed = {name for name in old_names if old_zip.read(name) != new_zip.read(name)}
        allowed_parts = {drawing_part, rel_path(drawing_part)}
        if allowed:
            allowed_parts.add(sheet_part)
        assert changed <= allowed_parts, changed - allowed_parts
        added = {name for name in new_names - old_names if not name.endswith("/")}
        assert len(added) == len(changed_slots) and all(name.startswith("xl/media/") for name in added), added
        if not allowed:
            assert old_zip.read(sheet_part) == new_zip.read(sheet_part), "Photo-only save changed worksheet bytes"
    semantic = {"last": str(after[title].cell(row, headings["Cal Last"]).value),
                "due": str(after[title].cell(row, headings["Cal Next"]).value),
                "smax": sha256(new_pictures[(row, headings["SMAX Tool Photo"])][0]["bytes"]).hexdigest(),
                "label": sha256(new_pictures[(row, headings["Label Photo"])][0]["bytes"]).hexdigest()}
    pictures = sum(len(sheet._images) for sheet in after)
    before.close()
    after.close()
    return {"cells_and_styles_preserved": checked, "pictures": pictures,
            "unselected_photos_and_zip_parts_unchanged": True,
            "changed_zip_parts": sorted(changed), "semantic": semantic}


def photos(temporary):
    paths = {}
    for kind, color in [("smax", (180, 220, 245)), ("calibration", (245, 220, 180))]:
        path = temporary / f"writer-{kind}.png"
        image = Image.new("RGB", (320, 200), color)
        ImageDraw.Draw(image).text((30, 80), kind.upper(), fill="black", font_size=30)
        image.save(path)
        paths[kind] = path
    return paths


def reordered_fixture(temporary, images, kind):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reordered"
    smax_heading = "Tool photo unavailable" if kind == "missing-column" else "SMAX Tool Photo"
    sheet.append(["Serial", smax_heading, "Cal Interval", "Cal Next", "Label Photo", "Asset #", "Cal Last", "Cal Cycle"])
    sheet.append([SERIAL, None, 6, datetime(2026, 6, 24), None, "South_TEST OWNER", datetime(2025, 12, 24), "MONTHS"])
    sheet["D2"].number_format = sheet["G2"].number_format = "yyyy-mm-dd"
    sheet.add_image(WorkbookImage(images["smax"]), "B2")
    sheet.add_image(WorkbookImage(images["calibration"]), "E2")
    if kind == "duplicate-smax":
        sheet.add_image(WorkbookImage(images["smax"]), "B2")
    if kind == "duplicate-label":
        sheet.add_image(WorkbookImage(images["calibration"]), "E2")
    workbook.create_sheet("Summary")["A1"] = "Preserve this worksheet"
    output = temporary / f"writer-{kind}.xlsx"
    workbook.save(output)
    workbook.close()
    return output


def install_inputs(page, images):
    page.evaluate("""()=>{for(const id of ['writer-smax','writer-calibration']){
        const input=document.createElement('input');input.type='file';input.id=id;input.hidden=true;document.body.append(input);
    }}""")
    for kind, path in images.items():
        page.locator(f"#writer-{kind}").set_input_files(path)


def writer_case(browser, url, workbook, images, temporary, scenario):
    context = browser.new_context(accept_downloads=True)
    context.add_init_script("(" + PICKER_SCRIPT + ")()")
    errors, remote = [], []
    context.on("request", lambda request: remote.append(request.url)
               if urlsplit(request.url).scheme in {"http", "https"}
               and urlsplit(request.url).netloc != urlsplit(url).netloc else None)
    page = context.new_page()
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(url)
    page.wait_for_function("window.ToolApp&&!ToolApp.state.busy&&!ToolApp.state.restoring")
    original_hash = seed_file(page, workbook)
    install_inputs(page, images)
    record = page.evaluate("""async({scenario,serial,last,due})=>{
        await ToolConnectedFile.connect();
        const snapshot=await ToolConnectedFile.read(),model=await ToolWorkbook.load(snapshot.bytes);
        const tool=model.tools.find(item=>item.serial===serial);
        const image=async kind=>({photoBytes:new Uint8Array(await document.querySelector('#writer-'+kind).files[0].arrayBuffer()),
            photoExtension:'png',photoWidth:320,photoHeight:200});
        const identity={toolId:tool.id,name:tool.name,serial:tool.serial};
        // SMAX can carry an unrelated last value. It never owns Cal Last.
        const smax={...identity,...await image('smax'),photoKind:'smax',updateDue:true,last:'2099-01-01',due};
        const calibration={...identity,...await image('calibration'),photoKind:'calibration',last,due};
        let updates;
        if(scenario==='photo-only')updates=[{...smax,updateDue:false,last:'2099-01-01',due:'not-a-date'}];
        else if(scenario==='photo-only-undefined'){
            const photoOnly={...smax,last:'2099-01-01',due:'not-a-date'};delete photoOnly.updateDue;updates=[photoOnly];
        }
        else if(scenario==='mixed-forward')updates=[calibration,smax];
        else if(scenario==='mixed-reverse')updates=[smax,calibration];
        else if(scenario==='mixed-photo-only')updates=[calibration,{...smax,updateDue:false,due:'9999-01-01'}];
        else if(scenario==='duplicate-smax-update')updates=[smax,{...smax}];
        else if(scenario==='conflicting-due')updates=[calibration,{...smax,due:'2027-01-22'}];
        else if(scenario==='legacy-calibration'){const legacy={...calibration};delete legacy.photoKind;updates=[legacy];}
        else if(scenario==='stale-due')updates=[{...smax,due:'2025-01-01'}];
        else if(scenario==='invalid-due')updates=[{...smax,due:'2026-02-30'}];
        else if(scenario==='invalid-kind')updates=[{...smax,photoKind:'unknown'}];
        else if(scenario==='duplicate-calibration-update')updates=[calibration,{...calibration}];
        else if(scenario==='duplicate-label-calibration')updates=[calibration];
        else updates=[smax];
        const validation=ToolWorkbook.validate(model,updates);
        let saved=false,applyError='';
        try{
            const bytes=await ToolWorkbook.apply(model,updates);
            await ToolConnectedFile.save({bytes,expectedVersion:snapshot.version});saved=true;
        }catch(error){applyError=error.message;}
        const backup=await ToolConnectedFile.takeBackup();
        return {validation,saved,applyError,tool:{photoCount:tool.photoCount,smaxPhotoCount:tool.smaxPhotoCount,
            hasSmaxPhotoColumn:tool.hasSmaxPhotoColumn,last:tool.last,due:tool.due},
            hash:await testHash(await(await testHandle.getFile()).arrayBuffer()),
            backupHash:backup?await testHash(backup.bytes):null,pickerCalls:testPickerCalls};
    }""", {"scenario": scenario, "serial": SERIAL, "last": LAST, "due": DUE})
    output = temporary / f"{workbook.stem}-{scenario}-result.xlsx"
    export_shared(page, output)
    assert not errors and not remote, (errors, remote)
    assert record["pickerCalls"] == 1, record
    assert record["tool"]["hasSmaxPhotoColumn"] is (scenario != "missing-column"), record
    assert record["tool"]["smaxPhotoCount"] == (0 if scenario == "missing-column" else 2 if scenario == "duplicate-smax" else 1), record
    assert record["tool"]["photoCount"] == (2 if "duplicate-label" in workbook.stem else 1), record
    blocked = scenario in {"duplicate-smax-update", "conflicting-due", "stale-due", "invalid-due", "invalid-kind",
                           "duplicate-calibration-update", "duplicate-label-calibration", "missing-column", "duplicate-smax"}
    if blocked:
        assert record["validation"] and record["applyError"] and not record["saved"], record
        assert record["hash"] == original_hash and output.read_bytes() == workbook.read_bytes(), record
        assert record["backupHash"] is None, record
        result = {"rejected_by_validate_and_apply": True, "real_opfs_byte_identical": True}
    else:
        assert not record["validation"] and not record["applyError"] and record["saved"], record
        assert record["hash"] != original_hash and record["backupHash"] == original_hash, record
        if scenario == "legacy-calibration":
            result = verify_output(workbook, output, None, images["calibration"], due=DUE, last=LAST)
        else:
            mixed = scenario in {"mixed-forward", "mixed-reverse", "mixed-photo-only"}
            result = verify_output(workbook, output, images["smax"], images["calibration"] if mixed else None,
                                   due=None if scenario.startswith("photo-only") else DUE,last=LAST if mixed else None)
        result.update({"real_opfs_saved": True, "exact_pre_write_backup": True})
    result["model"] = record["tool"]
    context.close()
    return result


def main():
    args = arguments(__doc__)
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix="tool-smax-workbook-") as directory, serve_program(args.program) as url, sync_playwright() as playwright:
        temporary = Path(directory)
        images = photos(temporary)
        browser = playwright.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
        cases = ["smax-due", "photo-only", "photo-only-undefined", "mixed-forward", "mixed-reverse", "mixed-photo-only", "legacy-calibration",
                 "duplicate-smax-update", "conflicting-due", "stale-due", "invalid-due", "invalid-kind", "duplicate-calibration-update"]
        result = {"company_workbook": {scenario: writer_case(browser, url, args.workbook, images, temporary, scenario)
                                      for scenario in cases}}
        result["reordered_headers"] = {}
        for fixture_kind in ["reordered", "missing-column", "duplicate-smax", "duplicate-label"]:
            fixture = reordered_fixture(temporary, images, fixture_kind)
            scenario = fixture_kind if fixture_kind in {"missing-column", "duplicate-smax"} else "smax-due"
            result["reordered_headers"][fixture_kind] = writer_case(browser, url, fixture, images, temporary, scenario)
            if fixture_kind == "duplicate-label":
                result["reordered_headers"]["duplicate-label-calibration"] = writer_case(
                    browser, url, fixture, images, temporary, "duplicate-label-calibration")
        assert result["company_workbook"]["mixed-forward"]["semantic"] == result["company_workbook"]["mixed-reverse"]["semantic"]
        browser.close()
    assert sha256(args.workbook.read_bytes()).hexdigest() == original_hash
    result["original_company_workbook_unchanged"] = True
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
