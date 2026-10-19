"""Real Chromium/XLSX tests. Inputs are read only; generated workbooks stay in /tmp.

python tests/browser_test.py [--program /path/Tool-Manager.html]
    [--workbook /absolute/path/company.xlsx] [--engine-only]
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime
from functools import partial
from hashlib import sha256
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import tempfile
from threading import Thread
from urllib.parse import quote
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from PIL import Image
from playwright.sync_api import sync_playwright

REPO = Path(__file__).resolve().parents[1]
LAST, DUE = "2026-12-29", "2027-12-29"


def arguments(description=__doc__):
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--photo", type=Path)
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--engine-only", action="store_true")
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    if args.workbook is None or not args.workbook.exists():
        parser.error("Supply --workbook /absolute/path/to/the/attached/company.xlsx")
    return args


@contextmanager
def serve_program(program):
    class QuietHandler(SimpleHTTPRequestHandler):
        def log_message(self, format, *args):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(program.resolve().parent)))
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}/{quote(program.name)}"
    finally:
        server.shutdown()
        thread.join()
        server.server_close()


def fixture_photo(args, temporary):
    source = args.photo
    if source is None:
        source = temporary / "original-label.png"
        with ZipFile(args.workbook) as archive:
            source.write_bytes(archive.read("xl/media/image45.png"))
    output = temporary / "test-label.png"
    with Image.open(source) as image:
        image = image.convert("RGB")
        image.putpixel((0, 0), (255, 0, 255))
        image.save(output)
    return output


def verify_workbook(original, updated, last=LAST, due=DUE):
    before, after = load_workbook(original), load_workbook(updated)
    assert before.sheetnames == after.sheetnames
    assert after["광주"]["A6"].value == datetime.fromisoformat(due)
    assert after["광주"]["B6"].value == datetime.fromisoformat(last)
    assert sum(len(s._images) for s in before) == sum(len(s._images) for s in after) == 59
    checked = 0
    for old_sheet in before:
        new_sheet = after[old_sheet.title]
        assert (old_sheet.max_row, old_sheet.max_column) == (new_sheet.max_row, new_sheet.max_column)
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)
        for row in old_sheet:
            for old_cell in row:
                new_cell = new_sheet[old_cell.coordinate]
                if (old_sheet.title, old_cell.coordinate) not in {("광주", "A6"), ("광주", "B6")}:
                    assert old_cell.value == new_cell.value, (old_sheet.title, old_cell.coordinate)
                assert old_cell._style == new_cell._style, (old_sheet.title, old_cell.coordinate)
                checked += 1
    allowed = {"xl/worksheets/sheet2.xml", "xl/drawings/drawing2.xml", "xl/drawings/_rels/drawing2.xml.rels", "[Content_Types].xml"}
    with ZipFile(original) as old_zip, ZipFile(updated) as new_zip:
        old_names, new_names = set(old_zip.namelist()), set(new_zip.namelist())
        assert old_names <= new_names, old_names - new_names
        changed = sorted(n for n in old_names if old_zip.read(n) != new_zip.read(n))
        assert not set(changed) - allowed, set(changed) - allowed
        added = {n for n in new_names - old_names if not n.endswith("/")}
        assert len(added) == 1 and all(n.startswith("xl/media/") for n in added), added
    before.close()
    after.close()
    return {"cells_checked": checked, "pictures": 59, "changed_zip_entries": changed}


def download_expression(page, expression, output):
    with page.expect_download() as event:
        page.evaluate("""async expression => {
            const bytes=await eval('(async()=>('+expression+'))()');
            const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([bytes]));
            a.download='test-result.xlsx';document.body.append(a);a.click();a.remove();
        }""", expression)
    event.value.save_as(output)


def engine_checks(page, args, temporary, photo):
    page.set_content('<input id="engine-workbook" type="file"><input id="engine-photo" type="file">')
    page.add_script_tag(path=REPO / "web/vendor/jszip.min.js")
    page.add_script_tag(path=REPO / "web/workbook.js")
    page.locator("#engine-workbook").set_input_files(args.workbook)
    page.locator("#engine-photo").set_input_files(photo)
    width, height = Image.open(photo).size
    checks = page.evaluate("""async ({width,height})=>{
      const model=await ToolWorkbook.load(await document.querySelector('#engine-workbook').files[0].arrayBuffer());
      const tool=model.tools.find(t=>t.sheet==='광주'&&t.row===6);
      const u={toolId:tool.id,name:'강건욱',serial:'31550075',last:'2026-12-29',due:'2027-12-29',photoBytes:new Uint8Array(await document.querySelector('#engine-photo').files[0].arrayBuffer()),photoExtension:'png',photoWidth:width,photoHeight:height};
      const v=x=>ToolWorkbook.validate(model,x),r={count:model.tools.length,tool,valid:v(u),name:v({...u,name:'다른사람'}),serial:v({...u,serial:'31550076'}),stale:v({...u,last:'2024-01-01'}),invalidDate:v({...u,last:'2026-02-30'}),reversed:v({...u,due:'2026-12-28'}),duplicate:v([u,u])};
      const disposed=model.tools.find(t=>t.due==='폐기'),ambiguous=model.tools.find(t=>t.photoCount>1);
      r.disposed=v({...u,toolId:disposed.id,name:disposed.name,serial:disposed.serial});
      r.ambiguous=v({...u,toolId:ambiguous.id,name:ambiguous.name,serial:ambiguous.serial});
      r.calendar=[ToolWorkbook.addMonths('2024-08-31',6),ToolWorkbook.addMonths('2024-02-29',12),ToolWorkbook.monthInterval('12','MONTHS')];
      window.engineModel=model;window.engineUpdate=u;return r;
    }""", {"width": width, "height": height})
    assert checks["count"] == 30
    assert checks["tool"]["name"] == "강건욱" and checks["tool"]["serial"] == "31550075"
    assert checks["tool"]["last"] == "2025-12-29" and checks["tool"]["due"] == "2026-12-29"
    assert checks["tool"]["months"] == 12 and not checks["valid"]
    for key in ["name", "serial", "stale", "invalidDate", "reversed", "duplicate", "disposed", "ambiguous"]:
        assert checks[key], f"Unsafe update accepted: {key}"
    assert checks["calendar"] == ["2025-02-28", "2025-02-28", 12]
    updated = temporary / "engine-result.xlsx"
    download_expression(page, "await ToolWorkbook.apply(engineModel,[engineUpdate])", updated)
    result = verify_workbook(args.workbook, updated)
    with ZipFile(updated) as archive:
        added_media = [n for n in archive.namelist() if n.startswith("xl/media/tool-label-")]
        assert len(added_media) == 1 and archive.read(added_media[0]) == photo.read_bytes()
    result["invalid_updates_blocked"] = 8
    result["fixtures"] = fixture_checks(page, temporary)
    return result


def fixture_checks(page, temporary):
    fixture = temporary / "fixture-1904.xlsx"
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.title = "Fixture"
    sheet.append(["Cal Next", "Cal Last", "Serial", "Asset #", "Cal Interval", "Cal Cycle", "Label Photo"])
    sheet.append([datetime(2025, 2, 28), datetime(2024, 2, 29), "001234", "South_Fixture_강건욱", "12", "MONTHS"])
    sheet["A2"].number_format = sheet["B2"].number_format = "yyyy-mm-dd"
    workbook.create_sheet("Summary")["A1"] = "Preserve this sheet"
    workbook.save(fixture)
    page.locator("#engine-workbook").set_input_files(fixture)
    result = page.evaluate("""async()=>{
      window.fixtureModel=await ToolWorkbook.load(await document.querySelector('#engine-workbook').files[0].arrayBuffer());
      const t=fixtureModel.tools[0];window.fixtureUpdate={...engineUpdate,toolId:t.id,serial:'001234',last:'2025-02-28',due:'2026-02-28'};
      return {date1904:fixtureModel.date1904,serial:t.serial,last:t.last,due:t.due,warnings:fixtureModel.warnings,valid:ToolWorkbook.validate(fixtureModel,fixtureUpdate),badZero:ToolWorkbook.validate(fixtureModel,{...fixtureUpdate,serial:'1234'})};
    }""")
    assert result["date1904"] and result["serial"] == "001234"
    assert result["last"] == "2024-02-29" and result["due"] == "2025-02-28"
    assert not result["valid"] and result["badZero"] and result["warnings"]
    output = temporary / "fixture-result.xlsx"
    download_expression(page, "await ToolWorkbook.apply(fixtureModel,[fixtureUpdate])", output)
    after = load_workbook(output)
    assert after.epoch == CALENDAR_MAC_1904 and after["Fixture"]["C2"].value == "001234"
    assert after["Fixture"]["A2"].value == datetime(2026, 2, 28)
    assert after["Fixture"]["B2"].value == datetime(2025, 2, 28)
    assert len(after["Fixture"]._images) == 1 and after["Summary"]["A1"].value == "Preserve this sheet"
    sheet.append([datetime(2025, 2, 28), datetime(2024, 2, 29), "001234", "South_Fixture_다른사람", "12", "MONTHS"])
    duplicate = temporary / "fixture-duplicate.xlsx"
    workbook.save(duplicate)
    page.locator("#engine-workbook").set_input_files(duplicate)
    assert page.evaluate("""async()=>{const m=await ToolWorkbook.load(await document.querySelector('#engine-workbook').files[0].arrayBuffer());return ToolWorkbook.validate(m,{...fixtureUpdate,toolId:m.tools[0].id});}""")
    workbook.close()
    after.close()
    return {"date1904": True, "new_drawing": True, "leading_zero_serial": True, "duplicate_serial_blocked": True}


def manual_checks(page, args, temporary, photo):
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.locator("#workbook").set_input_files(args.workbook)
    page.wait_for_function("window.ToolApp&&ToolApp.state.model&&!ToolApp.state.busy")
    assert page.evaluate("ToolApp.state.model.tools.length") == 30
    # New versions provide this switch; disable automatic recognition for the
    # manual review test. Automated flow has its own integration test.
    if page.locator("#automatic-mode").count():
        page.locator("#automatic-mode").uncheck()
    page.locator("#photos").set_input_files(photo)
    page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy")
    tool_id = page.evaluate("ToolApp.state.model.tools.find(t=>t.sheet==='광주'&&t.row===6).id")
    page.locator("#tool-select").select_option(tool_id)
    assert page.locator("#name").input_value() == "강건욱" and page.locator("#serial").input_value() == "31550075"
    page.locator("#last").fill(LAST)
    page.locator("#calculate-due").click()
    assert page.locator("#due").input_value() == DUE
    for field, value, message in [("name", "다른사람", "다릅니다"), ("due", "2026-12-28", "뒤여야"), ("last", "2024-01-01", "오래된")]:
        page.locator(f"#{field}").fill(value)
        page.locator("#confirm").check()
        page.locator("#save-review").click()
        assert message in page.locator("#status").inner_text()
        page.locator(f"#{field}").fill({"name": "강건욱", "due": DUE, "last": LAST}[field])
    page.locator("#confirm").check()
    page.locator("#save-review").click()
    page.wait_for_function("ToolApp.state.photos.filter(p=>p.review).length===1")
    page.locator("#serial").fill("31550076")
    assert not page.locator("#confirm").is_checked()
    assert page.evaluate("ToolApp.state.photos.filter(p=>p.review).length") == 0
    page.locator("#serial").fill("31550075")
    page.locator("#confirm").check()
    page.locator("#save-review").click()
    page.wait_for_function("!ToolApp.state.busy&&!ToolApp.state.preparing")
    with page.expect_download() as event:
        page.locator("#download-log").click()
    log = temporary / "changes.json"
    event.value.save_as(log)
    changes = json.loads(log.read_text())["changes"]
    assert len(changes) == 1 and changes[0]["serial"] == "31550075" and changes[0]["row"] == 6
    with page.expect_download() as event:
        page.locator("#export").click()
    output = temporary / "manual-result.xlsx"
    event.value.save_as(output)
    page.wait_for_function("!ToolApp.state.busy")
    result = verify_workbook(args.workbook, output)
    page.screenshot(path="/tmp/tool-manager-browser.png", full_page=True)
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate("document.documentElement.scrollWidth<=window.innerWidth"), "Mobile horizontal overflow"
    page.screenshot(path="/tmp/tool-manager-mobile.png", full_page=True)
    assert not errors, errors
    result.update({"approval_invalidated_on_edit": True, "change_log": True, "mobile_layout": True, "page_errors": errors})
    return result


def main():
    args = arguments()
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix="tool-browser-test-") as temporary_string:
        temporary = Path(temporary_string)
        photo = fixture_photo(args, temporary)
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
            context = browser.new_context(accept_downloads=True, locale="ko-KR")
            page = context.new_page()
            result = {"engine": engine_checks(page, args, temporary, photo)}
            if not args.engine_only:
                with serve_program(args.program) as url:
                    page.goto(url)
                    result["manual_ui"] = manual_checks(page, args, temporary, photo)
            browser.close()
    assert original_hash == sha256(args.workbook.read_bytes()).hexdigest(), "Original workbook changed"
    result["original_unchanged"] = True
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
