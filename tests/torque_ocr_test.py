"""Read the actual torque-tool evidence embedded in the company workbook.

The close-up calibration photo and the tool-system screenshot are different
original images. The calibration close-up must remain pending without identity;
the system screenshot may now update its SMAX slot without supplying Cal Last.
This does not reproduce the newly supplied inline-only photograph. Image bytes,
source drawing anchors, leading-zero serial, and expected worksheet dates are
checked without generating substitute text or replacing OCR responses.
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright

from browser_test import REPO, serve_program


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--program", type=Path, default=REPO / "web/dist/Tool-Manager.html")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--serial", default="0716805829")
    parser.add_argument("--chromium", default="/usr/bin/chromium")
    parser.add_argument("--diagnose-only", action="store_true")
    parser.add_argument("--diagnostics", type=Path, default=Path("/tmp/tool-torque-ocr.json"))
    args = parser.parse_args()
    if args.workbook is None:
        args.workbook = next(Path("/workspace/attachments").rglob("*.xlsx"), None)
    if args.workbook is None or not args.workbook.is_file():
        parser.error("Supply --workbook with the attached company XLSX.")
    if not args.program.is_file():
        parser.error("Supply --program with the standalone HTML program.")
    return args


def extract_sources(args, temporary):
    book = load_workbook(args.workbook)
    targets = []
    for sheet in book:
        headers = {str(cell.value): cell.column for cell in sheet[1]}
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, headers["Serial"]).value) == args.serial:
                targets.append((sheet, row, headers))
    assert len(targets) == 1, "The requested torque serial must identify exactly one original row"
    sheet, row, headers = targets[0]
    expected = {"sheet": sheet.title, "row": row, "serial": str(sheet.cell(row, headers["Serial"]).value), "name": str(sheet.cell(row, headers["Asset #"]).value).rsplit("_", 1)[-1].strip(), "last": sheet.cell(row, headers["Cal Last"]).value.strftime("%Y-%m-%d"), "due": sheet.cell(row, headers["Cal Next"]).value.strftime("%Y-%m-%d")}
    cases = []
    with ZipFile(args.workbook) as archive:
        originals = {sha256(archive.read(name)).hexdigest(): name for name in archive.namelist() if name.startswith("xl/media/") and not name.endswith("/")}
        for kind, column in [("calibration-label", headers["Label Photo"] - 1), ("system-screenshot", headers["Label Photo"])]:
            anchored = [image for image in sheet._images if image.anchor._from.row == row - 1 and image.anchor._from.col == column]
            assert len(anchored) == 1, (kind, row, column)
            data = anchored[0]._data()
            digest = sha256(data).hexdigest()
            media = originals[digest]
            assert data == archive.read(media)
            path = temporary / Path(media).name
            path.write_bytes(data)
            with Image.open(BytesIO(data)) as image:
                dimensions = list(image.size)
            cases.append({"kind": kind, "media": media, "sha256": digest, "dimensions": dimensions, "path": path})
    book.close()
    return expected, cases


def main():
    args = arguments()
    original_hash = sha256(args.workbook.read_bytes()).hexdigest()
    diagnostics = {"program": str(args.program), "diagnose_only": args.diagnose_only, "actual_original_images": True, "new_inline_photo_tested": False, "cases": []}
    failures = []
    with tempfile.TemporaryDirectory(prefix="tool-real-torque-") as directory:
        expected, cases = extract_sources(args, Path(directory))
        diagnostics["worksheet_reference"] = expected
        with serve_program(args.program) as url, sync_playwright() as p:
            browser = p.chromium.launch(executable_path=args.chromium, headless=True, args=["--no-sandbox"])
            context = browser.new_context(accept_downloads=True, locale="ko-KR")
            page = context.new_page()
            errors, requests = [], []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(url)
            page.wait_for_function("window.ToolApp&&window.ToolOCR&&window.ToolRuntime")
            # Observe raw real worker results without changing their contents.
            page.evaluate("""() => {
              window.torqueReadings=[];const create=Tesseract.createWorker.bind(Tesseract);
              Tesseract.createWorker=async(...args)=>{
                const worker=await create(...args),recognize=worker.recognize.bind(worker);
                worker.recognize=async(...input)=>{
                  const result=await recognize(...input),data=result.data||{};
                  const words=(data.words||[]).map(word=>({text:word.text,
                    confidence:word.confidence,bbox:word.bbox,
                    symbols:(word.symbols||[]).map(symbol=>({text:symbol.text,
                      confidence:symbol.confidence,bbox:symbol.bbox}))}));
                  torqueReadings.push({text:data.text,confidence:data.confidence,words});return result;
                };return worker;
              };
            }""")
            context.route("**/*", lambda route: (requests.append(route.request.url), route.abort()) if route.request.url.startswith(("http:", "https:")) else route.continue_())
            context.set_offline(True)
            page.locator("#workbook").set_input_files(args.workbook)
            page.wait_for_function("ToolApp.state.model&&!ToolApp.state.busy")
            tool_id = page.evaluate("serial=>ToolApp.state.model.tools.find(tool=>tool.serial===serial).id", args.serial)
            for case in cases:
                page.evaluate("torqueReadings=[]")
                page.locator("#photos").set_input_files(case["path"])
                page.wait_for_function("ToolApp.state.photos.length===1&&!ToolApp.state.busy&&!ToolApp.state.preparing", timeout=240000)
                evidence = page.evaluate("""() => {
                  const photo=ToolApp.state.photos[0];return {result:photo.result,
                    approval:photo.approval,assessment:photo.assessment,
                    prepared:!!ToolApp.state.preparedResult,readings:torqueReadings};
                }""")
                report = {key: value for key, value in case.items() if key != "path"}
                report.update(evidence)
                diagnostics["cases"].append(report)
                args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
                result = evidence["result"] or {}
                print(json.dumps({"kind": case["kind"], "media": case["media"], "actual_serial": result.get("serial"), "actual_last": result.get("last"), "actual_due": result.get("due"), "dates": result.get("dateEvidence"), "prepared": evidence["prepared"]}, ensure_ascii=False, indent=2), flush=True)
                if not args.diagnose_only:
                    if result.get("failed"):
                        failures.append(case["kind"] + ": real OCR did not finish")
                    if case["kind"] == "calibration-label":
                        if (result.get("last"), result.get("due")) != (expected["last"], expected["due"]):
                            failures.append("calibration-label: actual original dates were not read correctly")
                        if (result.get("dateEvidence") or {}).get("confidence", 0) is None or ((result.get("dateEvidence") or {}).get("confidence") or 0) < 70:
                            failures.append("calibration-label: the correct date pair lacks sufficient OCR confidence")
                        if "주식회사" in (result.get("identityEvidence") or {}).get("observedNames", []):
                            failures.append("calibration-label: company boilerplate was treated as an engineer name")
                        assert not evidence["prepared"], "A label with no engineer or tool serial must not select a tool"
                    else:
                        proof = [item for item in (result.get("smaxEvidence") or {}).get("serialReadings", [])
                                 if item.get("labelVerified") and item.get("value") == expected["serial"] and (item.get("confidence") or 0) >= 70]
                        assessment = evidence["assessment"] or {}
                        if result.get("documentType") != "smax" or not proof:
                            failures.append("system-screenshot: primary Serial/Asset Number was not read confidently")
                        assert result.get("last") == "", "SMAX must not invent an OCR Cal Last from audit dates"
                        assert evidence["prepared"] and assessment.get("eligible"), evidence
                        assert assessment.get("photoKind") == "smax" and assessment.get("toolId") == tool_id, assessment
                        assert assessment.get("serial") == expected["serial"], assessment
                        assert assessment.get("last") == expected["last"], "SMAX retains the workbook Cal Last"
                        assert assessment.get("due") == expected["due"], assessment
                page.locator("#remove-photo").click()
                page.wait_for_function("ToolApp.state.photos.length===0&&!ToolApp.state.preparing")
            assert not requests, requests
            assert not errors, errors
            browser.close()
    assert original_hash == sha256(args.workbook.read_bytes()).hexdigest(), "Original workbook changed"
    diagnostics.update({"failures": failures, "http_requests": requests, "page_errors": errors, "original_unchanged": True})
    args.diagnostics.write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"actual_original_cases": len(cases), "new_inline_photo_tested": False, "diagnose_only": args.diagnose_only, "failures": failures, "http_requests": requests, "original_unchanged": True, "diagnostics": str(args.diagnostics)}, ensure_ascii=False, indent=2))
    assert not failures, failures


if __name__ == "__main__":
    main()
